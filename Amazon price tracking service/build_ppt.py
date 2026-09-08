"""產生 Amazon Price Tracking Service 系統設計簡報（.pptx）。

從 DESIGN.md / README.md 的決策地圖、需求數學、API、資料模型、三個深入探討、
附錄（含 Amazon 反爬實測）整理成投影片。
執行：./.venv/bin/python build_ppt.py
"""
import math

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.oxml.ns import qn

# ---- palette ----
INK = RGBColor(0x0F, 0x17, 0x2A)
MUTED = RGBColor(0x64, 0x74, 0x8B)
ACCENT = RGBColor(0x25, 0x63, 0xEB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xF1, 0xF5, 0xF9)
CLIENT = RGBColor(0x64, 0x74, 0x8B)
GATEWAY = RGBColor(0x0E, 0xA5, 0xE9)
APP = RGBColor(0x25, 0x63, 0xEB)
DB = RGBColor(0x05, 0x96, 0x69)
CACHE = RGBColor(0xDC, 0x26, 0x26)
CDN = RGBColor(0x7C, 0x3A, 0xED)
ANALYTICS = RGBColor(0xD9, 0x77, 0x06)
GREEN = RGBColor(0x05, 0x96, 0x69)
RED = RGBColor(0xDC, 0x26, 0x26)
AMBER = RGBColor(0xD9, 0x77, 0x06)
RED_BG = RGBColor(0xFE, 0xE2, 0xE2)
GREEN_BG = RGBColor(0xDC, 0xFC, 0xE7)
CJK = "PingFang TC"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height


def _font(run, size, color, bold=False):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = CJK
    rPr = run._r.get_or_add_rPr()
    ea = rPr.makeelement(qn("a:ea"), {"typeface": CJK})
    rPr.append(ea)


def textbox(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, (text, size, color, bold, *rest) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(rest[0] if rest else 4)
        r = p.add_run()
        r.text = text
        _font(r, size, color, bold)
    return tb


def box(slide, x, y, w, h, text, fill, fg=WHITE, size=12, bold=True, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sp = slide.shapes.add_shape(shape, x, y, w, h)
    sp.fill.solid()
    sp.fill.fore_color.rgb = fill
    sp.line.color.rgb = fill
    sp.shadow.inherit = False
    tf = sp.text_frame
    tf.word_wrap = True
    tf.margin_top = Pt(2)
    tf.margin_bottom = Pt(2)
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = ln
        _font(r, size, fg, bold if i == 0 else False)
    return sp


def arrow(slide, x1, y1, x2, y2, color=MUTED, width=1.75):
    cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    cn.line.color.rgb = color
    cn.line.width = Pt(width)
    ln = cn.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return cn


def header(slide, kicker, title):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), SH)
    bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background(); bar.shadow.inherit = False
    textbox(slide, Inches(0.55), Inches(0.35), Inches(12), Inches(1.1),
            [(kicker, 13, ACCENT, True, 2), (title, 30, INK, True)])


def bullets(slide, x, y, w, h, items, size=15, gap=7):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    for i, it in enumerate(items):
        lvl = 0
        if isinstance(it, tuple):
            it, lvl = it
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.level = lvl
        r = p.add_run()
        prefix = "•  " if lvl == 0 else "–  "
        r.text = prefix + it
        _font(r, size if lvl == 0 else size - 1, INK if lvl == 0 else MUTED, False)
    return tb


def card(slide, x, y, w, h, fill=LIGHT):
    c = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    c.fill.solid(); c.fill.fore_color.rgb = fill; c.line.color.rgb = fill; c.shadow.inherit = False
    return c


def codebox(slide, x, y, w, h, code, size=11):
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = LIGHT; sp.line.color.rgb = LIGHT; sp.shadow.inherit = False
    tf = sp.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(12); tf.margin_right = Pt(6); tf.margin_top = Pt(8)
    for i, line in enumerate(code.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(1)
        r = p.add_run(); r.text = line if line else " "
        r.font.size = Pt(size); r.font.name = "Menlo"
        r.font.color.rgb = MUTED if line.strip().startswith("#") or line.strip().startswith("--") else INK
    return sp


def table(slide, rows, x, y, w, h, widths, size=12, head_size=13, first_col_accent=True, cell_fill=None):
    """rows[0] 是表頭。widths 為每欄 inch。cell_fill(r, c, text) 可覆寫底色。"""
    ncol = len(rows[0])
    tbl = slide.shapes.add_table(len(rows), ncol, x, y, w, h).table
    for i, cw in enumerate(widths):
        tbl.columns[i].width = Inches(cw)
    for r in range(len(rows)):
        for c in range(ncol):
            cell = tbl.cell(r, c)
            cell.margin_top = Pt(2); cell.margin_bottom = Pt(2); cell.margin_left = Pt(8); cell.margin_right = Pt(6)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            fill = None
            if cell_fill and r:
                fill = cell_fill(r, c, rows[r][c])
            cell.fill.fore_color.rgb = ACCENT if r == 0 else (fill or (WHITE if r % 2 else LIGHT))
            p = cell.text_frame.paragraphs[0]
            run = p.add_run(); run.text = rows[r][c]
            color = WHITE if r == 0 else (ACCENT if (c == 0 and first_col_accent) else INK)
            _font(run, size if r else head_size, color, r == 0 or c == 0)
    return tbl


def strip_rows(slide, items, y0=Inches(1.85), rh=Inches(1.0), gap=Inches(0.2), title_size=15, body_size=12.5):
    """左側色條 + 淺底卡片的列（面試框架 / 扣分點慣用）。items = [(title, body, strip_color)]"""
    cy = y0
    for title, body, strip in items:
        st = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), cy, Inches(0.12), rh)
        st.fill.solid(); st.fill.fore_color.rgb = strip; st.line.fill.background(); st.shadow.inherit = False
        bg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.72), cy, Inches(12.0), rh)
        bg.fill.solid(); bg.fill.fore_color.rgb = LIGHT; bg.line.fill.background(); bg.shadow.inherit = False
        textbox(slide, Inches(0.95), cy + Inches(0.1), Inches(11.6), rh - Inches(0.15),
                [(title, title_size, INK, True, 4), (body, body_size, MUTED, False)])
        cy += rh + gap


# ============ Slide 1: Title ============
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
bg.fill.solid(); bg.fill.fore_color.rgb = INK; bg.line.fill.background(); bg.shadow.inherit = False
box(s, Inches(0.9), Inches(2.0), Inches(1.5), Inches(1.5), "$↓", ACCENT, WHITE, 44)
textbox(s, Inches(0.9), Inches(3.7), Inches(11.5), Inches(2.5), [
    ("Amazon Price Tracking Service", 44, WHITE, True, 6),
    ("System Design demo — 價格歷史、降價通知與眾包爬取", 20, RGBColor(0xCB, 0xD5, 0xE1), False, 10),
    ("FastAPI 原型 · 真抓 amazon.com（curl_cffi TLS 指紋）· 對齊 PDF 設計題 · 附 AWS Terraform", 14, MUTED, False),
])

# ============ Slide 2: Requirements & scale ============
s = prs.slides.add_slide(BLANK)
header(s, "01 · REQUIREMENTS", "需求與規模")
cols = [
    ("功能性 (FR)", APP, [
        "查看 Amazon 商品價格歷史（網站 / Chrome extension）",
        "訂閱降價通知，自訂價格門檻",
        "Out of scope：站內搜尋、跨零售商比價、評論整合",
    ]),
    ("非功能性 (NFR)", DB, [
        "可用性 > 一致性（eventual consistency 可接受）",
        "5 億商品",
        "價格歷史查詢 < 500ms",
        "價格變動後 1 小時內通知",
    ]),
    ("容量數學", ANALYTICS, [
        "盲爬：5e8 / (1000 IP × 1 req/s) / 86400 ≈ 5.8 天掃一輪 → 淘汰",
        "寫入：5 億 × 每天 1 次 ≈ 5.8K writes/s → append-only、依商品分區",
        "歷史：每小時一筆 × 2 年 = 17,520 列 / 商品",
        "訂閱查詢：索引 (product_id, price_threshold, user_id) range scan",
    ]),
]
cw = Inches(3.9); gap = Inches(0.25); x0 = Inches(0.6); y0 = Inches(1.8)
for i, (title, color, items) in enumerate(cols):
    x = x0 + i * (cw + gap)
    box(s, x, y0, cw, Inches(0.6), title, color, WHITE, 16)
    card(s, x, y0 + Inches(0.7), cw, Inches(4.4))
    bullets(s, x + Inches(0.2), y0 + Inches(0.9), cw - Inches(0.4), Inches(4.0), items, size=13, gap=10)

# ============ Slide 3: API ============
s = prs.slides.add_slide(BLANK)
header(s, "02 · API DESIGN", "API 設計（對齊 PDF 兩支 + 眾包 / crawler / 彙總端點）")
rows = [
    ("方法", "路徑", "說明"),
    ("GET", "/api/v1/price/{product_id}?period=&granularity=", "PDF API 1：價格歷史 → {granularity, source: aggregation|raw_fallback, points[], rows_scanned, latency_ms, current_price}"),
    ("POST", "/api/v1/subscriptions", "PDF API 2：{user_id, product_id, price_threshold, notification_type: sse|email, email?}（upsert）"),
    ("POST", "/api/v1/price-reports", "extension 回報 {product_id, price, reporter_token} → accepted / pending_verify；429 超過限流"),
    ("GET", "/api/v1/price-reports?product_id=", "回報狀態機（pending_verify → confirmed / rejected）"),
    ("POST", "/api/v1/products/track", "{product_id}（ASIN 或 URL）→ 註冊 + 最高優先爬取"),
    ("GET / POST", "/api/v1/products · /products/{asin}/crawl?sync=", "商品 + 優先分數 / 間隔 / 到期 / fetcher / 狀態；立即爬取"),
    ("POST / GET", "/api/v1/aggregations/run · /aggregations/stats", "立即彙總；raw 列數 vs daily / weekly / monthly 列數"),
    ("GET", "/api/v1/stream?user_id= · /notifications?user_id=", "SSE 即時通知（dev-only）；outbox 狀態機"),
    ("GET", "/api/v1/cdc/status · /crawl/log", "CDC checkpoint / head / lag / queue 深度；爬取紀錄"),
    ("POST", "/api/v1/debug/seed · /debug/mock-price", "只在 MOCK_AMAZON=1（或 ALLOW_SEED=1）開放"),
]
table(s, rows, Inches(0.6), Inches(1.65), Inches(12.1), Inches(5.0), [1.3, 4.3, 6.5], size=11, head_size=12)
textbox(s, Inches(0.6), Inches(6.75), Inches(12.1), Inches(0.5), [
    ("身分：正式版 Cognito JWT（env-gated）；原型 user_id 由前端產生（一個分頁 = 一個使用者）。", 12, MUTED, False)])

# ============ Slide 4: Data model ============
s = prs.slides.add_slide(BLANK)
header(s, "03 · DATA MODEL", "資料模型（8 張表，時間一律 UTC）")
rows = [
    ("表", "重點欄位", "索引 / 備註"),
    ("products", "asin(PK) · title · currency · status(active|unavailable|region_locked|not_found|blocked) · last_price/_at/_source · last_seen_at · last_fetcher · subscription_count · view_count · suspicious_flag", "last_event_price / last_event_seq 只由 notify worker 更新（消費端與寫入端分離才能算 Δ）"),
    ("prices（append-only）", "seq（自增，模擬 WAL 位置）· product_id · price · currency · source(crawler|extension|mock|seed) · fetcher · ts", "INDEX(product_id, ts)；正式版 DynamoDB：hash product_id、range ts、Streams NEW_IMAGE"),
    ("subscriptions", "subscription_id(PK) · user_id · product_id · price_threshold · notification_type · email · status · last_notified_price（cooldown）", "UNIQUE(product_id, user_id)（PDF PK）+ INDEX(product_id, price_threshold, user_id)（PDF secondary，covering）"),
    ("price_reports", "report_id · product_id · reported_price · reporter_token（匿名）· baseline_price · deviation_pct · verify_price", "status: accepted | pending_verify | confirmed | rejected"),
    ("price_aggregations", "avg / min / max / close_price · sample_count", "PK(product_id, granularity, bucket_start)；Postgres，讀路徑只碰這張"),
    ("notification_outbox", "event_seq · old_price · new_price · threshold · channel", "notification_id = sha256(subscription_id|event_seq)[:32]；status: ENQUEUED|ATTEMPTED|VENDOR_ACCEPTED|FAILED|SKIPPED_COOLDOWN"),
    ("cdc_checkpoint", "consumer(prices_seq / ddb:{shard_id}) · position", "at-least-once：發完事件才存 checkpoint"),
    ("crawl_log", "asin · fetcher · status · elapsed_ms · ts", "demo 面板用；cleanup cron 清"),
]
table(s, rows, Inches(0.6), Inches(1.6), Inches(12.15), Inches(5.3), [2.1, 5.7, 4.35], size=11.5, head_size=12.5)

# ============ Slide 5: High-level architecture ============
s = prs.slides.add_slide(BLANK)
header(s, "04 · HIGH-LEVEL ARCHITECTURE", "整體架構")
# row A: read/write path
yA = Inches(1.65); bh = Inches(0.8)
box(s, Inches(0.5), yA, Inches(1.7), bh, "Client / extension\n網站 · Chrome", CLIENT, WHITE, 11)
box(s, Inches(2.5), yA, Inches(1.5), bh, "API Gateway\nJWT · 節流", GATEWAY, WHITE, 11)
box(s, Inches(4.3), yA, Inches(2.2), bh, "API (FastAPI)\n/price · /subscriptions\n/price-reports · /products", APP, WHITE, 11)
box(s, Inches(7.0), yA, Inches(2.3), bh, "PostgreSQL (RDS)\nproducts · subscriptions\naggregations · outbox", DB, WHITE, 11)
box(s, Inches(9.8), yA, Inches(3.0), bh, "prices：DynamoDB\n(PK product_id, SK ts) append-only\n+ DynamoDB Streams (CDC)", DB, WHITE, 11)
arrow(s, Inches(2.2), yA + Inches(0.4), Inches(2.5), yA + Inches(0.4))
arrow(s, Inches(4.0), yA + Inches(0.4), Inches(4.3), yA + Inches(0.4))
arrow(s, Inches(6.5), yA + Inches(0.4), Inches(7.0), yA + Inches(0.4))
arrow(s, Inches(6.5), yA + Inches(0.6), Inches(9.8), yA + Inches(0.6), color=GREEN, width=1.2)
# row B: crawler pipeline
yB = Inches(3.05)
box(s, Inches(0.5), yB, Inches(1.7), bh, "Scheduler\n優先分數 · due 排序", ANALYTICS, WHITE, 11)
box(s, Inches(2.5), yB, Inches(1.5), bh, "crawl queue\n(Redis Streams)", CACHE, WHITE, 11)
box(s, Inches(4.3), yB, Inches(2.2), bh, "Crawl worker\ncurl_cffi → Playwright\ntoken bucket 1 rps", ANALYTICS, WHITE, 11)
box(s, Inches(7.0), yB, Inches(2.3), bh, "amazon.com\n(美國出口 IP / proxy pool)", INK, WHITE, 11)
arrow(s, Inches(2.2), yB + Inches(0.4), Inches(2.5), yB + Inches(0.4))
arrow(s, Inches(4.0), yB + Inches(0.4), Inches(4.3), yB + Inches(0.4))
arrow(s, Inches(6.5), yB + Inches(0.4), Inches(7.0), yB + Inches(0.4))
arrow(s, Inches(5.4), yB, Inches(5.4), yA + bh, color=GREEN, width=1.2)  # worker → API/prices（寫價）
textbox(s, Inches(5.5), yA + bh + Inches(0.05), Inches(1.6), Inches(0.35), [("append price", 10, GREEN, False)])
arrow(s, Inches(1.35), yA + bh, Inches(1.35), yB, color=MUTED, width=1.0)  # extension report → scheduler(urgent)
textbox(s, Inches(1.4), yA + bh + Inches(0.05), Inches(1.4), Inches(0.35), [("可疑 → urgent", 10, MUTED, False)])
# row C: CDC → notify
yC = Inches(4.45)
box(s, Inches(9.8), yC, Inches(3.0), bh, "CDC tailer\nper-shard checkpoint", CDN, WHITE, 11)
box(s, Inches(7.0), yC, Inches(2.3), bh, "price_changed\n(Redis Streams)", CACHE, WHITE, 11)
box(s, Inches(4.3), yC, Inches(2.2), bh, "price-change worker\njoin 訂閱 · 過濾 · outbox", CDN, WHITE, 11)
box(s, Inches(2.5), yC, Inches(1.5), bh, "notify:*\n(sse / email)", CACHE, WHITE, 11)
box(s, Inches(0.5), yC, Inches(1.7), bh, "Sender\nSES email / SSE", GREEN, WHITE, 11)
arrow(s, Inches(11.3), yA + bh, Inches(11.3), yC, color=CDN)
arrow(s, Inches(9.8), yC + Inches(0.4), Inches(9.3), yC + Inches(0.4))
arrow(s, Inches(7.0), yC + Inches(0.4), Inches(6.5), yC + Inches(0.4))
arrow(s, Inches(4.3), yC + Inches(0.4), Inches(4.0), yC + Inches(0.4))
arrow(s, Inches(2.5), yC + Inches(0.4), Inches(2.2), yC + Inches(0.4))
# row D: aggregator
yD = Inches(5.75)
box(s, Inches(9.8), yD, Inches(3.0), Inches(0.7), "Aggregator job\n每 5 分鐘增量 / 夜間 EventBridge", ANALYTICS, WHITE, 11)
box(s, Inches(7.0), yD, Inches(2.3), Inches(0.7), "price_aggregations\ndaily / weekly / monthly", DB, WHITE, 11)
arrow(s, Inches(9.8), yD + Inches(0.35), Inches(9.3), yD + Inches(0.35))
arrow(s, Inches(11.3), yC + bh, Inches(11.3), yD, color=MUTED, width=1.0)
textbox(s, Inches(0.5), Inches(5.75), Inches(6.2), Inches(1.4), [
    ("三條路徑：", 12, INK, True, 2),
    ("① 資料從哪來 — extension 眾包 + 優先式 crawler，寫入端只 append prices", 11.5, MUTED, False, 2),
    ("② 變動怎麼變通知 — CDC tail Streams → worker join 訂閱 → outbox → sender", 11.5, MUTED, False, 2),
    ("③ 歷史怎麼快速讀 — aggregator 預彙總，/price 讀 30 列而非 17,520 列", 11.5, MUTED, False)])

# ============ Slide 6: Decision map ============
s = prs.slides.add_slide(BLANK)
header(s, "決策地圖", "三大關鍵架構決策（環環相扣）")
cards = [
    ("① 如何追蹤 5 億商品", APP, "extension 眾包 + 優先式爬取\n+ 完整性驗證", [
        "盲爬 1000 IP 掃一輪 5.8 天 → 太舊",
        "誰重要誰先爬：訂閱數 / 查看數當分數",
        "extension 回報使用者正在看的價格",
        "代價：使用者資料不能直接信 → 可疑不入庫、最高優先重爬",
    ]),
    ("② 如何 1 小時內通知", DB, "CDC → Redis Streams\n→ price-change worker", [
        "cron 每 2h 全表掃：延遲吃頻率、每次 full scan",
        "寫入端只 append；CDC tail 變更日誌變事件",
        "worker 用 (product_id, threshold, user_id) 索引 join",
        "代價：at-least-once → outbox 冪等 + cooldown",
    ]),
    ("③ 如何 < 500ms 畫圖", CACHE, "pre-aggregation 表\n(daily / weekly / monthly)", [
        "每小時一筆 × 2 年 = 17,520 列，date_trunc+avg 撐不住",
        "背景 job 預先彙總，30 天圖讀 30 列、2 年讀 24 列",
        "代價：新鮮度落後一個週期；即時價另外給",
        "新商品 raw fallback 照樣能畫",
    ]),
]
cw = Inches(3.9); gp = Inches(0.3); x0 = Inches(0.5); y = Inches(1.75)
for i, (title, color, choice, items) in enumerate(cards):
    x = x0 + i * (cw + gp)
    box(s, x, y, cw, Inches(0.6), title, color, WHITE, 15)
    box(s, x, y + Inches(0.68), cw, Inches(0.75), choice, GREEN, WHITE, 12)
    card(s, x, y + Inches(1.52), cw, Inches(3.1))
    bullets(s, x + Inches(0.2), y + Inches(1.68), cw - Inches(0.4), Inches(2.8), items, size=12, gap=8)
    if i < 2:
        arrow(s, x + cw, y + Inches(0.3), x + cw + gp, y + Inches(0.3))
textbox(s, Inches(0.5), Inches(6.5), Inches(12.3), Inches(0.6), [
    ("①決定「資料從哪來」、②決定「變動怎麼變成通知」、③決定「歷史怎麼快速讀」—— 下面逐項展開。", 13, MUTED, False)])

# ============ Slide 7: Deep dive 1 — evolution ============
s = prs.slides.add_slide(BLANK)
header(s, "深入 1 · DISCOVERY", "發現並追蹤 5 億商品：三階段演進")
stages = [
    ("(Naive) 盲爬", CLIENT, "seed → BFS 抓連結 → 平行 → 去重", [
        "Amazon 每 IP 限 1 req/s",
        "1000 IP 也要 5.8 天掃一輪",
        "資料太舊、絕大多數商品沒人在意",
    ]),
    ("(Better) 優先式爬取", GATEWAY, "誰重要誰先爬", [
        "訊號：訂閱數（有人在等）、查看數（有人在看）",
        "流量集中在少數商品 → 效率大增",
        "冷啟動：新熱門商品沒人搜過就沒分數",
    ]),
    ("(Best) extension 眾包", APP, "使用者逛 Amazon，extension 回報 (asin, 價格)", [
        "天然涵蓋「有人在意」的商品、即時",
        "順便發現新商品；crawler 只補近期沒人看的",
        "代價：上傳資料不能直接信 → 完整性驗證",
    ]),
]
cw = Inches(3.9); gp = Inches(0.3); x0 = Inches(0.5); y = Inches(1.85)
for i, (title, color, sub, items) in enumerate(stages):
    x = x0 + i * (cw + gp)
    box(s, x, y, cw, Inches(0.6), title, color, WHITE, 15)
    box(s, x, y + Inches(0.68), cw, Inches(0.6), sub, LIGHT, INK, 12, bold=False)
    card(s, x, y + Inches(1.36), cw, Inches(2.6))
    bullets(s, x + Inches(0.2), y + Inches(1.5), cw - Inches(0.4), Inches(2.4), items, size=12.5, gap=8)
    if i < 2:
        arrow(s, x + cw, y + Inches(0.3), x + cw + gp, y + Inches(0.3))
box(s, Inches(0.5), Inches(6.05), Inches(12.3), Inches(0.85),
    "把最大的限制（要監控上億商品）變成優勢：使用者的瀏覽行為就是最好的優先訊號。\ncrawler 結果永遠是真相，extension 只負責「告訴系統該看哪裡」。", DB, WHITE, 13)

# ============ Slide 8: Deep dive 1 — priority formulas ============
s = prs.slides.add_slide(BLANK)
header(s, "深入 1 · SCHEDULER", "優先分數、爬取間隔與 due-ness（app/scheduler.py）")
codebox(s, Inches(0.6), Inches(1.65), Inches(7.3), Inches(2.3), """score    = 10·subscription_count + 2·log1p(view_count)
         + 100·suspicious_flag + priority_boost
interval = clamp(3600 / (1 + score), 60s, 86400s)
# 分數越高爬越勤；captcha 狀態 ×4 退避；not_found 拉到最大間隔
due      = now - last_seen_at >= interval
# extension 回報也算「看過」→ crawler 可以晚點再來
rank     = score × age / interval      # 每 tick 取 due 商品排序取前 N""", size=12)
card(s, Inches(8.2), Inches(1.65), Inches(4.55), Inches(2.3))
bullets(s, Inches(8.4), Inches(1.8), Inches(4.2), Inches(2.1), [
    "log1p(view_count)：查看數邊際遞減，避免被刷",
    "衰減：用 age / interval 排序，久沒爬的自然浮上來，冷門商品不餓死（上限 24h）",
    "訂閱 → +10 分；可疑回報 → +100 分",
], size=12, gap=7)
textbox(s, Inches(0.6), Inches(4.15), Inches(12.2), Inches(0.4), [("排程與插隊", 15, INK, True)])
steps = [
    ("scheduler tick", "取 due 商品\n依 rank 取前 N", ANALYTICS),
    ("crawl queue", "FIFO channel\n(Redis Streams)", CACHE),
    ("crawl worker", "token bucket 1 rps\ncurl_cffi → Playwright", APP),
    ("append prices", "source=crawler\n→ CDC → 通知", DB),
]
n = len(steps); bw = Inches(2.5); gp = Inches(0.3); x0 = Inches(0.6); y = Inches(4.6)
for i, (t, d, color) in enumerate(steps):
    x = x0 + i * (bw + gp)
    box(s, x, y, bw, Inches(1.1), t + "\n" + d, color, WHITE, 12)
    if i < n - 1:
        arrow(s, x + bw, y + Inches(0.55), x + bw + gp, y + Inches(0.55))
box(s, Inches(0.6), Inches(5.9), Inches(5.3), Inches(0.8), "enqueue_urgent()\n訂閱新商品 / 手動爬 / 可疑回報 → 直接入列 + 設 flag", RED, WHITE, 12)
arrow(s, Inches(5.9), Inches(6.3), Inches(6.5), Inches(5.7), color=RED, width=1.2)
textbox(s, Inches(6.7), Inches(5.95), Inches(6.1), Inches(0.9), [
    ("job 掉了下個 tick 會撿回（flag 還在）。正式版：Redis ZSET 當 priority queue、多 crawler consumer group、per-IP token bucket 集中 Redis、proxy pool 輪替（附錄 C）。", 12, MUTED, False)])

# ============ Slide 9: Deep dive 1 — integrity verification flow ============
s = prs.slides.add_slide(BLANK)
header(s, "深入 1 · INTEGRITY", "完整性驗證：使用者上傳的資料不能直接信（附錄 D）")
box(s, Inches(0.5), Inches(1.75), Inches(2.4), Inches(0.9), "extension 回報\n(asin, price, reporter_token)", CLIENT, WHITE, 11)
box(s, Inches(3.3), Inches(1.75), Inches(3.0), Inches(0.9), "baseline = products.last_price\ndeviation = (baseline − reported) / baseline", APP, WHITE, 11)
arrow(s, Inches(2.9), Inches(2.2), Inches(3.3), Inches(2.2))
# accepted branch
box(s, Inches(3.3), Inches(3.35), Inches(3.0), Inches(0.9), "跌幅 ≤ 30% → accepted\nappend(source=extension) → CDC → 通知\nlast_seen_at 更新", GREEN, WHITE, 11)
arrow(s, Inches(4.8), Inches(2.65), Inches(4.8), Inches(3.35), color=GREEN)
# suspicious branch
box(s, Inches(6.9), Inches(1.75), Inches(3.0), Inches(0.9), "沒 baseline 或跌幅 > 30%\n→ pending_verify（不寫 prices 表）", AMBER, WHITE, 11)
arrow(s, Inches(6.3), Inches(2.2), Inches(6.9), Inches(2.2), color=AMBER)
box(s, Inches(10.3), Inches(1.75), Inches(2.5), Inches(0.9), "suspicious_flag=1\nenqueue_urgent 最高優先重爬", RED, WHITE, 11)
arrow(s, Inches(9.9), Inches(2.2), Inches(10.3), Inches(2.2), color=RED)
box(s, Inches(10.3), Inches(3.35), Inches(2.5), Inches(0.9), "crawler 回來\n|crawl − reported| / crawl", ANALYTICS, WHITE, 11)
arrow(s, Inches(11.55), Inches(2.65), Inches(11.55), Inches(3.35), color=RED)
box(s, Inches(6.9), Inches(3.35), Inches(3.0), Inches(0.9), "≤ 5% → confirmed", GREEN, WHITE, 12)
box(s, Inches(6.9), Inches(4.45), Inches(3.0), Inches(0.9), "> 5% → rejected", RED, WHITE, 12)
arrow(s, Inches(10.3), Inches(3.7), Inches(9.9), Inches(3.7), color=GREEN)
arrow(s, Inches(10.3), Inches(4.0), Inches(9.9), Inches(4.8), color=RED)
box(s, Inches(3.3), Inches(4.45), Inches(3.0), Inches(0.9), "入庫的永遠是 crawler 價\n→ CDC → 通知", DB, WHITE, 11)
arrow(s, Inches(6.9), Inches(3.9), Inches(6.3), Inches(4.8), color=GREEN, width=1.2)
bullets(s, Inches(0.5), Inches(5.6), Inches(12.3), Inches(1.5), [
    "可疑回報不寫 prices 表 → 不會觸發通知；惡意回報只會讓系統多爬一次（reporter_token 每分鐘 30 次限流），不會造成假通知",
    "容忍度 5% 留給幣別換算 / 優惠券差異；匿名：不綁 user_id、不記瀏覽紀錄（附錄 I，隱私友善）",
    "demo 面板：「正常回報 −3%」直接入庫；「可疑回報 −60%」→ pending_verify → 重爬 → confirmed / rejected",
], size=12.5, gap=6)

# ============ Slide 10: Amazon anti-bot field notes ============
s = prs.slides.add_slide(BLANK)
header(s, "深入 1 · 反爬實測", "真抓 amazon.com：踩過的雷（附錄 E，2026-09）")
notes = [
    ("TLS 指紋", "Amazon 在 TLS handshake 就比對 JA3/HTTP2 指紋，純 requests 直接被擋 → curl_cffi（curl-impersonate）讓指紋與真瀏覽器一致", CACHE),
    ("impersonate=chrome124 被攔", "拿到「Click the button below to continue shopping」攔截頁（3.7KB、HTTP 200）；chrome120 / safari17_0 / edge101 拿到完整頁（1.2–2.7MB）→ 預設 chrome120，被擋就輪替指紋 + 新 cookie jar", AMBER),
    ("幣別 cookie", "i18n-prefs=USD 固定顯示幣別，否則依 IP 換算，且 .a-offscreen 格式不穩", GATEWAY),
    ("主價格區塊解析", "只從 #corePrice_feature_div / #apex_desktop … 的第一個 .a-price 組價（whole + fraction）；整頁第一個 .a-price 會抓到相關商品輪播；「Save 20% with Trade-In」會被當 $20", APP),
    ("region_locked", "非美國 IP 看到「cannot be shipped to your selected delivery location」→ buybox 沒價格；改配送地要登入 → 判成 region_locked（非缺貨）；正式版要美國出口 IP / proxy pool", CDN),
    ("Playwright fallback 成本", "能跑 JS、較能過 captcha，但映像 +400MB、每次吃 CPU/記憶體 → 只當 fallback（CRAWLER_PLAYWRIGHT=1）；限速 token bucket 1 rps + jitter", ANALYTICS),
    ("robots / ToS", "robots.txt 對 /dp/ 無明確 disallow 但 ToS 禁自動化；本專案為教學 demo：1 rps、不登入、不繞 captcha；正式產品應評估 Product Advertising API / 合法資料供應商", MUTED),
]
strip_rows(s, notes, y0=Inches(1.55), rh=Inches(0.74), gap=Inches(0.04), title_size=13, body_size=12)

# ============ Slide 11: Deep dive 2 — cron vs event-driven ============
s = prs.slides.add_slide(BLANK)
header(s, "深入 2 · NOTIFY", "處理價格變動並通知：cron 全表掃 → event-driven")
box(s, Inches(0.6), Inches(1.75), Inches(5.9), Inches(0.6), "現況：cron 每 2 小時掃 price 表", RED, WHITE, 15)
card(s, Inches(0.6), Inches(2.45), Inches(5.9), Inches(2.4))
bullets(s, Inches(0.8), Inches(2.6), Inches(5.5), Inches(2.2), [
    "延遲吃 cron 頻率：最壞 2 小時才發現 → 違反 1 小時 SLA",
    "每次昂貴 full scan：5 億商品 × 每次都掃",
    "縮短 cron 間隔只是把 full scan 跑得更頻繁",
    "pull 批次 → 改成 push 單一事件",
], size=13, gap=8)
box(s, Inches(6.85), Inches(1.75), Inches(5.9), Inches(0.6), "方案：CDC → Redis Streams → worker", GREEN, WHITE, 15)
card(s, Inches(6.85), Inches(2.45), Inches(5.9), Inches(2.4))
bullets(s, Inches(7.05), Inches(2.6), Inches(5.5), Inches(2.2), [
    "寫入端只 append price 表，不必記得發事件",
    "CDC tail 變更日誌，每筆新價格 → 一個事件（< 1s）",
    "worker 每事件只查該商品的訂閱（covering index）",
    "at-least-once → outbox 冪等 + cooldown 兜底",
], size=13, gap=8)
arrow(s, Inches(6.5), Inches(2.05), Inches(6.85), Inches(2.05))
steps = [
    ("append prices", "crawler / extension\n→ DynamoDB", DB),
    ("DynamoDB Streams", "變更日誌\n(原型：seq 輪詢)", CDN),
    ("CDC tailer", "checkpoint\nat-least-once", CDN),
    ("price_changed", "Redis Streams\nper-channel", CACHE),
    ("price-change worker", "stream ⋈ subscriptions\n過濾 · outbox", APP),
    ("sender", "SES email / SSE\nVENDOR_ACCEPTED", GREEN),
]
n = len(steps); bw = Inches(1.85); gp = Inches(0.2); x0 = Inches(0.6); y = Inches(5.15)
for i, (t, d, color) in enumerate(steps):
    x = x0 + i * (bw + gp)
    box(s, x, y, bw, Inches(1.2), t + "\n" + d, color, WHITE, 11)
    if i < n - 1:
        arrow(s, x + bw, y + Inches(0.6), x + bw + gp, y + Inches(0.6))
textbox(s, Inches(0.6), Inches(6.5), Inches(12.2), Inches(0.5), [
    ("回填不是變動：/debug/seed 灌歷史後把 checkpoint 推過回填列，否則 17,520 筆假事件會炸通知。", 12, MUTED, False)])

# ============ Slide 12: CDC vs dual-write ============
s = prs.slides.add_slide(BLANK)
header(s, "深入 2 · CDC vs DUAL-WRITE", "價格事件怎麼產生：CDC vs dual-write（附錄 A）")
rows = [
    ("面向", "CDC（log-based，本專案）", "Dual-write"),
    ("一致性", "寫入端只寫 DB；事件從 log 衍生，不會漏", "DB 成功、發事件失敗 → 漏通知；反過來 → 幽靈通知。需 transactional outbox 補救"),
    ("順序", "log 順序 = commit 順序（DynamoDB Streams per-shard 有序）", "併發寫入時事件順序不保證"),
    ("過濾 / 合併", "在 consumer 端做（worker 過濾 < 1% 波動）", "可在寫入端就做"),
    ("侵入性", "寫入端零改動；多一個 tailer process", "每個寫入路徑都要記得發事件"),
    ("延遲", "log → 事件通常 < 1s（DynamoDB Streams 近即時）", "即時"),
]
table(s, rows, Inches(0.6), Inches(1.7), Inches(12.15), Inches(3.6), [1.8, 5.0, 5.35], size=12.5)
box(s, Inches(0.6), Inches(5.6), Inches(12.15), Inches(1.2),
    "本專案：CDC。原型 SQLite/Postgres 的 seq 輪詢是「窮人版 log tail」（SELECT … WHERE seq > checkpoint ORDER BY seq LIMIT 500，每秒一次）；\n正式版換 DynamoDB Streams，事件格式不變、consumer 不用改。", DB, WHITE, 13)

# ============ Slide 13: DynamoDB Streams + Redis vs Kafka ============
s = prs.slides.add_slide(BLANK)
header(s, "深入 2 · DYNAMODB STREAMS", "DynamoDB Streams 機制與三種消費形態（附錄 B）")
card(s, Inches(0.6), Inches(1.65), Inches(5.9), Inches(2.3))
textbox(s, Inches(0.8), Inches(1.75), Inches(5.5), Inches(0.4), [("Streams 機制", 14, INK, True)])
bullets(s, Inches(0.8), Inches(2.15), Inches(5.5), Inches(1.8), [
    "記錄 24 小時內每筆 INSERT / MODIFY / REMOVE，依 partition 分 shard",
    "同一 shard 內有序；同一 item 的變更一定在同一 shard 鏈",
    "describe_stream → get_shard_iterator(TRIM_HORIZON | AFTER_SEQUENCE_NUMBER) → get_records 迴圈",
    "shard 會關閉並分裂成 child shard → 要定期重列",
], size=11.5, gap=5)
card(s, Inches(6.85), Inches(1.65), Inches(5.9), Inches(2.3))
textbox(s, Inches(7.05), Inches(1.75), Inches(5.5), Inches(0.4), [("Redis Streams vs Kafka", 14, INK, True)])
bullets(s, Inches(7.05), Inches(2.15), Inches(5.5), Inches(1.8), [
    "1 小時 SLA、每秒數千事件 → Redis Streams 夠用，compose 輕、與 Earthquake 同一套",
    "per-channel 好切：crawl / price_changed / notify:sse / notify:email",
    "要長期回放 / 多下游（OLAP、data warehouse）→ Kafka 當 log，Redis 當工作佇列",
    "SQS 淘汰：沒 consumer group",
], size=11.5, gap=5)
rows = [
    ("消費形態", "做法", "優", "劣"),
    ("① 自己 tail（本專案）", "worker 用 boto3 直接讀；checkpoint（SequenceNumber）存 Postgres cdc_checkpoint", "單一 codebase、不需 Lambda", "多 worker 時要自己做 shard lease"),
    ("② Lambda event source mapping", "AWS 管 checkpoint / 重試；Lambda 把事件轉丟 Redis Streams", "免管 checkpoint", "多一個 runtime"),
    ("③ KCL + Kinesis Adapter", "DynamoDB Streams Kinesis Adapter + KCL", "多 worker shard lease 現成", "較重"),
]
table(s, rows, Inches(0.6), Inches(4.15), Inches(12.15), Inches(2.6), [2.6, 4.6, 2.3, 2.65], size=11.5, head_size=12)

# ============ Slide 14: price-change worker + outbox state machine ============
s = prs.slides.add_slide(BLANK)
header(s, "深入 2 · WORKER", "price-change worker + outbox 狀態機（附錄 G）")
steps = [
    ("① prev", "prev = last_event_price\n（消費端自己記，與寫入端分離）", CDN),
    ("② join 索引", "WHERE product_id AND status='active'\nAND price_threshold >= new", DB),
    ("③ 小波動過濾", "|Δ| < 1% 且沒人「這次才跨過門檻」\n→ 不通知、不寫 outbox", AMBER),
    ("④ cooldown", "已通知 $120 → $118 再通知\n$119.5 不通知 (SKIPPED_COOLDOWN)", GATEWAY),
    ("⑤ outbox 冪等", "id = sha256(sub_id|event_seq)\n重播 / 重啟不重送", APP),
    ("⑥ 回升重設", "價格回到門檻之上\n→ last_notified_price = NULL", GREEN),
]
n = len(steps); bw = Inches(1.95); gp = Inches(0.12); x0 = Inches(0.5); y = Inches(1.7)
for i, (t, d, color) in enumerate(steps):
    x = x0 + i * (bw + gp)
    box(s, x, y, bw, Inches(1.35), t + "\n" + d, color, WHITE, 12)
    if i < n - 1:
        arrow(s, x + bw, y + Inches(0.67), x + bw + gp, y + Inches(0.67))
textbox(s, Inches(0.5), Inches(3.3), Inches(12.2), Inches(0.4), [("outbox 狀態機（sender per channel）", 15, INK, True)])
states = [
    ("ENQUEUED", "worker 寫 outbox\n丟 notify:{type}", APP),
    ("ATTEMPTED", "sender terminal 檢查\n→ 標記送出中", GATEWAY),
    ("VENDOR_ACCEPTED", "SES / SSE 成功\n寫 last_notified_price", GREEN),
    ("FAILED", "vendor 拒絕 / 例外\n可觀測、可重試", RED),
    ("SKIPPED_COOLDOWN", "沒更便宜 → 不送\n仍記一筆讓狀態可觀測", AMBER),
]
bw = Inches(2.3); gp = Inches(0.17); x0 = Inches(0.5); y = Inches(3.8)
for i, (t, d, color) in enumerate(states):
    x = x0 + i * (bw + gp)
    box(s, x, y, bw, Inches(1.1), t + "\n" + d, color, WHITE, 11)
arrow(s, x0 + bw, y + Inches(0.55), x0 + bw + gp, y + Inches(0.55))
arrow(s, x0 + 2 * bw + gp, y + Inches(0.55), x0 + 2 * (bw + gp), y + Inches(0.55), color=GREEN)
arrow(s, x0 + 2 * bw + gp, y + Inches(0.85), x0 + 3 * (bw + gp), y + Inches(1.0), color=RED, width=1.2)
bullets(s, Inches(0.5), Inches(5.2), Inches(12.3), Inches(1.9), [
    "本質是 event stream ⋈ DB：每個事件只查該商品的訂閱，走 (product_id, price_threshold, user_id) covering index，不掃全表",
    "at-least-once（tailer 重播、worker 重啟）由 notification_id 冪等吸收：同一價格事件對同一訂閱只會有一筆 outbox",
    "dual-write 論點裡的「寫入端過濾 / 合併」在 CDC 架構下改放 consumer 端做（第 ③ 步）",
    "demo 已驗證：降價通知（SSE）→ 同價過濾 → 再跌再通知；outbox 面板每 3 秒輪詢看狀態機變化",
], size=12.5, gap=6)

# ============ Slide 15: Deep dive 3 — pre-aggregation ============
s = prs.slides.add_slide(BLANK)
header(s, "深入 3 · HISTORY", "價格歷史圖表：raw 查詢 → pre-aggregation")
codebox(s, Inches(0.6), Inches(1.65), Inches(5.9), Inches(1.3), """SELECT date_trunc('day', ts), avg(price) FROM prices
 WHERE product_id = :pid AND ts >= now() - interval '2y'
 GROUP BY 1
-- 熱門商品掃 17,520 列；數百萬使用者同時畫圖撐不住""", size=11)
box(s, Inches(6.85), Inches(1.65), Inches(5.9), Inches(1.3),
    "pre-aggregation（app/aggregation.py）\n背景 job 把 raw 彙總進 price_aggregations\nPK=(product_id, granularity, bucket_start) · avg/min/max/close/n", APP, WHITE, 12)
arrow(s, Inches(6.5), Inches(2.3), Inches(6.85), Inches(2.3))
rows = [
    ("period", "粒度", "彙總列數（每小時一筆）", "raw 列數"),
    ("7d / 30d / 90d", "daily", "7 / 30 / 90", "168 / 720 / 2,160"),
    ("1y", "weekly", "52", "8,760"),
    ("2y", "monthly", "24", "17,520"),
]
table(s, rows, Inches(0.6), Inches(3.2), Inches(5.9), Inches(1.7), [1.6, 1.2, 1.7, 1.4], size=12.5)
rows2 = [
    ("實測（本機 SQLite）", "rows_scanned", "latency"),
    ("30 天圖 · raw fallback", "720 列", "2.9 ms"),
    ("30 天圖 · 彙總後", "31 列", "0.65 ms"),
    ("2 年圖 · 彙總後", "25 列", "0.57 ms"),
]
table(s, rows2, Inches(6.85), Inches(3.2), Inches(5.9), Inches(1.7), [2.7, 1.6, 1.6], size=12.5,
      cell_fill=lambda r, c, t: GREEN_BG if r > 1 else None)
bullets(s, Inches(0.6), Inches(5.15), Inches(12.2), Inches(1.9), [
    "API 依 period 選粒度：≤ 90d daily、≤ 1y weekly、> 1y monthly → 30 天圖讀 30 列、2 年圖讀 24 列，毫秒級回",
    "17,520 列彙總後 → daily 731 / weekly 106 / monthly 25（最新 bucket 是「開放中」的，每次彙總會重算）",
    "原型每 5 分鐘增量重算有新價格的商品 + 「立即彙總」按鈕；正式版夜間 EventBridge job（可掛 read replica 專門掃）",
    "儲存線性成長：每商品 daily 365 × 年 + weekly + monthly；DynamoDB 只放 raw，彙總表在 Postgres",
], size=12.5, gap=6)

# ============ Slide 16: Freshness, raw fallback, OLAP ============
s = prs.slides.add_slide(BLANK)
header(s, "深入 3 · TRADE-OFFS", "新鮮度、raw fallback 與 OLAP / TSDB 替代（附錄 F）")
cols = [
    ("新鮮度", AMBER, [
        "彙總最多落後一個週期（正式版 ≤ 24h）",
        "圖表看趨勢可接受",
        "即時價從 products.last_price 另外給（current_price 欄位）",
        "最新 bucket 開放中，每次彙總重算",
    ]),
    ("raw fallback", GATEWAY, [
        "新商品還沒彙總 → 現算 date_trunc + avg",
        "回應標 source=\"raw_fallback\"，圖照樣能畫",
        "demo 面板顯示 raw 列數 vs 彙總列數與延遲，看得到差異",
        "彙總 job 掛掉 → 讀路徑仍可畫圖（fail-soft，附錄 M）",
    ]),
    ("OLAP / TSDB 替代", CDN, [
        "ClickHouse / InfluxDB（column-oriented）對 avg/min/max 很快",
        "適合跨商品分析、heavy join",
        "ingestion：OLTP → CDC → stream → OLAP",
        "本專案的 CDC 管線加一個 consumer 即可灌 OLAP",
    ]),
]
cw = Inches(3.9); gap = Inches(0.25); x0 = Inches(0.6); y0 = Inches(1.75)
for i, (title, color, items) in enumerate(cols):
    x = x0 + i * (cw + gap)
    box(s, x, y0, cw, Inches(0.6), title, color, WHITE, 16)
    card(s, x, y0 + Inches(0.7), cw, Inches(3.5))
    bullets(s, x + Inches(0.2), y0 + Inches(0.9), cw - Inches(0.4), Inches(3.2), items, size=13, gap=10)
box(s, Inches(0.6), Inches(6.1), Inches(12.2), Inches(0.8),
    "取捨：用「落後一個週期」換「讀 30 列 vs 17,520 列」；歷史圖不需要即時，即時價另給 —— 兩條讀路徑各司其職。", DB, WHITE, 13)

# ============ Slide 17: Prototype vs Production ============
s = prs.slides.add_slide(BLANK)
header(s, "PROTOTYPE → PRODUCTION", "原型 vs Production（全 env-gated，本機零依賴）")
rows = [
    ("面向", "原型（本機）", "Production（雲端）", "切換 env"),
    ("Price DB", "SQLite prices（seq）", "DynamoDB（product_id, ts）", "PRICE_TABLE"),
    ("CDC", "seq 輪詢 tailer", "DynamoDB Streams tailer", "PRICE_TABLE"),
    ("Queue", "asyncio.Queue（per-channel）", "Redis Streams", "REDIS_URL"),
    ("目錄 / 訂閱 / 彙總 / outbox", "SQLite", "RDS PostgreSQL", "DATABASE_URL"),
    ("Crawler", "真抓 amazon.com（curl_cffi）；MOCK_AMAZON=1 離線", "同，worker ASG + 美國出口 IP / proxy pool", "MOCK_AMAZON · AMAZON_DOMAIN · CRAWL_RPS"),
    ("Playwright fallback", "未裝不啟用", "--build-arg WITH_PLAYWRIGHT=1", "CRAWLER_PLAYWRIGHT"),
    ("通知", "SSE 到瀏覽器；email dry-run", "SES email", "SES_FROM_EMAIL"),
    ("彙總", "in-proc 每 5 分鐘", "EventBridge 夜間 job + in-proc 增量", "AGG_INTERVAL_S"),
    ("Scheduler / tailer / workers", "主程序 asyncio task", "worker ASG（ROLE=worker）", "ROLE"),
    ("reporter 限流", "記憶體 deque", "Redis INCR + EXPIRE", "部署形態"),
    ("Auth", "無（前端產 user_id）", "Cognito JWT", "AUTH_ENABLED"),
]
table(s, rows, Inches(0.6), Inches(1.6), Inches(12.15), Inches(4.9), [2.6, 3.6, 3.5, 2.45], size=11.5, head_size=12)
textbox(s, Inches(0.6), Inches(6.6), Inches(12.2), Inches(0.5), [
    ("沿用 QR / Earthquake 的「工廠函式 + 延遲匯入」寫法：未設 env → 記憶體 / SQLite 實作；docker compose 走 Postgres + Redis Streams 的 production code path。", 12, MUTED, False)])

# ============ Slide 18: AWS deployment + CI/CD ============
s = prs.slides.add_slide(BLANK)
header(s, "部署 · AWS", "AWS 部署架構（Terraform，ap-northeast-1）+ CI/CD")
yA = Inches(1.6); bh = Inches(0.8)
box(s, Inches(0.5), yA, Inches(1.6), bh, "Client /\nextension", CLIENT, WHITE, 11)
box(s, Inches(2.3), yA, Inches(1.7), bh, "CloudFront\n+ WAF per-IP", CDN, WHITE, 11)
box(s, Inches(4.2), yA, Inches(1.7), bh, "API Gateway\nJWT · 節流", GATEWAY, WHITE, 11)
box(s, Inches(6.1), yA, Inches(1.4), bh, "內部 ALB\n/health", APP, WHITE, 11)
box(s, Inches(7.7), yA, Inches(1.8), bh, "EC2 ASG (api)\nDocker/gunicorn", INK, WHITE, 11)
box(s, Inches(4.2), yA + Inches(1.0), Inches(1.7), Inches(0.6), "Cognito (JWT)", DB, WHITE, 11)
arrow(s, Inches(2.1), yA + Inches(0.4), Inches(2.3), yA + Inches(0.4))
arrow(s, Inches(4.0), yA + Inches(0.4), Inches(4.2), yA + Inches(0.4))
arrow(s, Inches(5.9), yA + Inches(0.4), Inches(6.1), yA + Inches(0.4))
arrow(s, Inches(7.5), yA + Inches(0.4), Inches(7.7), yA + Inches(0.4))
arrow(s, Inches(5.05), yA + Inches(1.0), Inches(5.05), yA + bh, color=MUTED, width=1.0)
textbox(s, Inches(5.95), yA + Inches(0.75), Inches(1.2), Inches(0.3), [("VPC Link", 10, MUTED, False)])
# stores
sx = Inches(9.9); sw = Inches(2.9)
box(s, sx, Inches(1.5), sw, Inches(0.7), "RDS PostgreSQL\nproducts · subscriptions · aggregations · outbox", DB, WHITE, 11)
box(s, sx, Inches(2.3), sw, Inches(0.7), "DynamoDB prices\nappend-only · Streams NEW_IMAGE = CDC", DB, WHITE, 11)
box(s, sx, Inches(3.1), sw, Inches(0.7), "ElastiCache Redis\nStreams: crawl / price_changed / notify:*", CACHE, WHITE, 11)
for ty in (Inches(1.85), Inches(2.65), Inches(3.45)):
    arrow(s, Inches(9.5), yA + Inches(0.4), sx, ty, color=MUTED, width=1.2)
# worker ASG
box(s, Inches(2.3), Inches(3.3), Inches(5.2), Inches(0.85), "EC2 ASG (worker)  ROLE=worker\ncrawler(→ amazon.com) · CDC tailer(DynamoDB Streams) · notify workers", ANALYTICS, WHITE, 11)
box(s, Inches(7.7), Inches(3.3), Inches(1.8), Inches(0.85), "SES\n(email)", GREEN, WHITE, 11)
arrow(s, Inches(7.5), Inches(3.72), Inches(7.7), Inches(3.72), color=GREEN)
arrow(s, Inches(7.5), Inches(3.5), sx, Inches(2.65), color=MUTED, width=1.0)
arrow(s, Inches(7.5), Inches(3.95), sx, Inches(3.45), color=MUTED, width=1.0)
textbox(s, Inches(0.5), Inches(4.3), Inches(12.3), Inches(0.5), [
    ("Cognito(JWT) · SSM(設定) · Secrets Manager(DB 密碼) · ECR(映像) · EventBridge + Lambda(cleanup) · CloudWatch/SNS(監控告警)。SSE 過 CloudFront/API GW 有 buffering → dev-only。", 11, MUTED, False)])
textbox(s, Inches(0.5), Inches(4.85), Inches(12.2), Inches(0.4), [("CI/CD（GitHub Actions，deploy-price.yml，workflow_dispatch 手動觸發）", 14, INK, True)])
steps = [
    ("workflow_dispatch", "手動觸發\n(infra 尚未 apply)", CLIENT),
    ("GitHub Actions", "OIDC assume role\n(無長期金鑰)", GATEWAY),
    ("buildx", "--platform\nlinux/arm64", APP),
    ("ECR push", "tag = SHA\n+ latest", DB),
    ("SSM 更新", "/price/\nIMAGE_URI", CDN),
    ("SSM 部署", "send-command\ntag:app=price", ANALYTICS),
    ("上線", "deploy-app.sh\nALB /health", GREEN),
]
n = len(steps); bw = Inches(1.62); gp = Inches(0.13); x0 = Inches(0.5); y = Inches(5.3)
for i, (t, d, color) in enumerate(steps):
    x = x0 + i * (bw + gp)
    box(s, x, y, bw, Inches(1.2), t + "\n" + d, color, WHITE, 11)
    if i < n - 1:
        arrow(s, x + bw, y + Inches(0.6), x + bw + gp, y + Inches(0.6))
textbox(s, Inches(0.5), Inches(6.6), Inches(12.2), Inches(0.4), [
    ("前置 repo variable AWS_DEPLOY_ROLE_ARN = terraform output gha_deploy_role_arn；terraform apply 會計費，不用時 destroy。", 11, MUTED, False)])

# ============ Slide 19: Load test + monitoring ============
s = prs.slides.add_slide(BLANK)
header(s, "壓測 · 觀測", "k6 壓測三場景 + 監控指標（附錄 K）")
box(s, Inches(0.6), Inches(1.7), Inches(5.9), Inches(0.6), "k6 壓測（loadtest/）", APP, WHITE, 15)
rows = [
    ("場景", "路徑", "驗證"),
    ("history", "GET /price/{id}?period=", "p95 < 500ms（NFR）；aggregation vs raw_fallback"),
    ("report", "POST /price-reports", "extension 寫入路徑；限流 429、pending_verify 比例"),
    ("subscribe", "POST /subscriptions", "upsert 寫入；UNIQUE(product_id, user_id)"),
]
table(s, rows, Inches(0.6), Inches(2.4), Inches(5.9), Inches(2.0), [1.3, 2.1, 2.5], size=11.5, head_size=12)
box(s, Inches(6.85), Inches(1.7), Inches(5.9), Inches(0.6), "監控指標（app 層）", DB, WHITE, 15)
card(s, Inches(6.85), Inches(2.4), Inches(5.9), Inches(2.0))
bullets(s, Inches(7.05), Inches(2.5), Inches(5.5), Inches(1.9), [
    "CDC lag：checkpoint vs head（/cdc/status）、queue 深度",
    "crawl 成功率 / captcha 率 / region_locked 率 / 用了哪層 fetcher",
    "通知 p95 延遲：事件 ts → VENDOR_ACCEPTED（對 1 小時 SLA）",
    "raw_fallback 比例、/price latency_ms",
], size=12, gap=5)
box(s, Inches(0.6), Inches(4.65), Inches(12.15), Inches(0.5), "基礎設施層：CloudWatch alarms → SNS（ALB / RDS / Redis / DynamoDB / API GW / CloudFront）+ Synthetic canary 探 /health", CDN, WHITE, 12)
textbox(s, Inches(0.6), Inches(5.35), Inches(12.2), Inches(0.4), [("降級 / fail-soft（附錄 M）", 14, INK, True)])
bullets(s, Inches(0.6), Inches(5.75), Inches(12.2), Inches(1.4), [
    "crawler 被擋只影響該商品（退避 ×4），API 照常服務歷史；Redis 故障 → 事件停在 checkpoint 之後不會丟，tailer 重連後補發、冪等吸收",
    "彙總 job 掛掉 → 讀路徑 raw fallback 仍可畫圖；全域例外處理（500 不洩漏、422 可讀）、Request ID、body 大小上限 413",
    "Cleanup（附錄 L）：EventBridge + Lambda 每日清 outbox / crawl_log / 已結案 price_reports；prices / aggregations 是產品本身，不清",
], size=12, gap=6)

# ============ Slide 20: Interview framework + pitfalls ============
s = prs.slides.add_slide(BLANK)
header(s, "面試實戰 · FRAMEWORK", "45 分鐘怎麼講 + 常見扣分點")
iframe = [
    ("0–5 min：需求釐清", "只追 Amazon？要 extension 嗎？5 億商品、< 500ms、1 小時通知 → 先算 5.8 天那條數學，說明盲爬為何不可行", AMBER),
    ("5–15 min：API + 資料模型 + High-Level", "兩支 API、prices append-only（DynamoDB）+ subscriptions 索引 + 彙總表；畫出 crawler / CDC / aggregator 三條路徑", ACCENT),
    ("15–35 min：三個深入（面試官挑 1–2 追問）", "① 眾包 + 優先式爬取 + 完整性驗證｜② CDC → stream → worker（冪等 / cooldown）｜③ pre-aggregation + raw fallback", DB),
    ("35–45 min：取捨與收尾", "CDC vs dual-write、Redis vs Kafka、彙總新鮮度、反爬 / 合規；反問面試官", GREEN),
]
strip_rows(s, iframe, y0=Inches(1.55), rh=Inches(0.72), gap=Inches(0.1), title_size=13, body_size=11)
textbox(s, Inches(0.6), Inches(4.85), Inches(12.2), Inches(0.4), [("常見扣分點", 15, RED, True)])
pitfalls = [
    ("直接信使用者回報", "extension 回報沒驗證就入庫 → 假通知；要說「可疑不入庫、先重爬」"),
    ("cron 掃表找變動", "延遲吃 cron 頻率 + full scan；要說 event-driven / CDC"),
    ("raw 資料直接畫圖", "17,520 列 date_trunc；要說 pre-aggregation + 粒度選擇"),
    ("dual-write 沒 outbox", "DB 成功、事件失敗漏通知；at-least-once 又沒冪等會重送"),
    ("忽略反爬 / 合規", "沒提 TLS 指紋、限速、region、ToS；沒算 5 億 × 1 rps 的可行性"),
]
cw = Inches(2.35); gp = Inches(0.11); x0 = Inches(0.6); y = Inches(5.3)
for i, (t, d) in enumerate(pitfalls):
    x = x0 + i * (cw + gp)
    st = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(0.1), Inches(1.5))
    st.fill.solid(); st.fill.fore_color.rgb = RED; st.line.fill.background(); st.shadow.inherit = False
    card(s, x + Inches(0.1), y, cw - Inches(0.1), Inches(1.5))
    textbox(s, x + Inches(0.2), y + Inches(0.08), cw - Inches(0.3), Inches(1.4),
            [(t, 12.5, INK, True, 3), (d, 11, MUTED, False)])

# ============ Slide 21: Appendix overview ============
s = prs.slides.add_slide(BLANK)
header(s, "附錄總覽", "DESIGN.md 附錄 A–M 一頁列表")
apps = [
    ("A", "CDC vs dual-write", "一致性 / 順序 / 過濾 / 侵入性 / 延遲比較；本專案 CDC"),
    ("B", "DynamoDB Streams 機制與消費", "shard 有序、iterator 迴圈；自己 tail / Lambda ESM / KCL；Redis Streams vs Kafka"),
    ("C", "優先式爬取：分數、衰減、升級", "log1p 邊際遞減、age/interval 排序、captcha 退避；正式版 Redis ZSET + consumer group"),
    ("D", "完整性驗證流程", "跌幅 > 30% → pending_verify → 重爬 → ≤ 5% confirmed；惡意回報只多爬一次"),
    ("E", "Amazon 反爬實測筆記", "TLS 指紋、chrome124 攔截頁、幣別 cookie、主價格區塊、region_locked、Playwright、ToS"),
    ("F", "彙總粒度與 OLAP 替代", "period → 粒度 → 列數表；新鮮度；ClickHouse / InfluxDB 走 CDC 多接 consumer"),
    ("G", "冪等、cooldown 與回升重設", "notification_id = hash(sub|seq)；$120 → $118 才再通知；回到門檻上重設"),
    ("H", "SSE dev-only vs SES email", "SSE 過 CDN 有 buffering；SES 需驗證身分、bounce/complaint 回饋"),
    ("I", "Auth 與匿名回報", "Cognito JWT env-gated；extension 回報只帶 reporter_token，不記瀏覽紀錄"),
    ("J", "Rate limiting", "WAF per-IP、API GW 節流、reporter_token 每分鐘 30 次、出站 CRAWL_RPS"),
    ("K", "Monitoring", "CloudWatch → SNS、canary；CDC lag、captcha 率、通知 p95、queue 深度"),
    ("L", "Cleanup cron", "EventBridge + Lambda 每日清 outbox / crawl_log / 結案回報；prices 不清"),
    ("M", "Error handling / 降級", "500 不洩漏、422 可讀、Request ID、413；crawler / Redis / 彙總各自 fail-soft"),
]
col_w = Inches(6.0); x0 = Inches(0.6); y0 = Inches(1.6); rh = Inches(0.76)
for i, (letter, title, desc) in enumerate(apps):
    col, row = (0, i) if i < 7 else (1, i - 7)
    x = x0 + col * (col_w + Inches(0.2)); y = y0 + row * rh
    box(s, x, y + Inches(0.05), Inches(0.5), Inches(0.6), letter, ACCENT, WHITE, 16)
    textbox(s, x + Inches(0.6), y, col_w - Inches(0.7), rh,
            [(title, 13, INK, True, 1), (desc, 11.5, MUTED, False)])

prs.save("Amazon_Price_Tracking_Service.pptx")
print("saved Amazon_Price_Tracking_Service.pptx ·", len(prs.slides._sldIdLst), "slides")
