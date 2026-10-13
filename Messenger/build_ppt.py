"""產生 Messenger 系統設計簡報（.pptx）。

從 DESIGN.md / README.md 的決策地圖、需求數學、協定比較、API、資料模型、三個深入探討、
failover / fencing 時序、實測結果與部署整理成投影片。
執行：./.venv/bin/python build_ppt.py
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.oxml.ns import qn

# ---- palette ----
INK = RGBColor(0x0F, 0x17, 0x2A)
MUTED = RGBColor(0x64, 0x74, 0x8B)
ACCENT = RGBColor(0x25, 0x63, 0xEB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xF1, 0xF5, 0xF9)
CLIENT = RGBColor(0x64, 0x74, 0x8B)
GATEWAY = RGBColor(0x0E, 0xA5, 0xE9)
APP = RGBColor(0x25, 0x63, 0xEB)
DB = RGBColor(0x05, 0x96, 0x69)
CACHE = RGBColor(0xDC, 0x26, 0x26)
CDN = RGBColor(0x7C, 0x3A, 0xED)
ANALYTICS = RGBColor(0xD9, 0x77, 0x06)
GREEN = RGBColor(0x05, 0x96, 0x69)
RED = RGBColor(0xDC, 0x26, 0x26)
AMBER = RGBColor(0xD9, 0x77, 0x06)
RED_BG = RGBColor(0xFE, 0xE2, 0xE2)
GREEN_BG = RGBColor(0xDC, 0xFC, 0xE7)
CJK = "PingFang TC"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height


def _font(run, size, color, bold=False):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = CJK
    rPr = run._r.get_or_add_rPr()
    ea = rPr.makeelement(qn("a:ea"), {"typeface": CJK})
    rPr.append(ea)


def textbox(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, (text, size, color, bold, *rest) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(rest[0] if rest else 4)
        r = p.add_run()
        r.text = text
        _font(r, size, color, bold)
    return tb


def box(slide, x, y, w, h, text, fill, fg=WHITE, size=12, bold=True, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sp = slide.shapes.add_shape(shape, x, y, w, h)
    sp.fill.solid()
    sp.fill.fore_color.rgb = fill
    sp.line.color.rgb = fill
    sp.shadow.inherit = False
    tf = sp.text_frame
    tf.word_wrap = True
    tf.margin_top = Pt(2)
    tf.margin_bottom = Pt(2)
    tf.margin_left = Pt(3)
    tf.margin_right = Pt(3)
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = ln
        _font(r, size, fg, bold if i == 0 else False)
    return sp


def line(slide, x1, y1, x2, y2, color=MUTED, width=1.75):
    cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    cn.line.color.rgb = color
    cn.line.width = Pt(width)
    return cn


def arrow(slide, x1, y1, x2, y2, color=MUTED, width=1.75):
    cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    cn.line.color.rgb = color
    cn.line.width = Pt(width)
    ln = cn.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return cn


def header(slide, kicker, title):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), SH)
    bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background(); bar.shadow.inherit = False
    textbox(slide, Inches(0.55), Inches(0.35), Inches(12), Inches(1.1),
            [(kicker, 13, ACCENT, True, 2), (title, 30, INK, True)])


def bullets(slide, x, y, w, h, items, size=15, gap=7):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    for i, it in enumerate(items):
        lvl = 0
        if isinstance(it, tuple):
            it, lvl = it
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.level = lvl
        r = p.add_run()
        prefix = "•  " if lvl == 0 else "–  "
        r.text = prefix + it
        _font(r, size if lvl == 0 else size - 1, INK if lvl == 0 else MUTED, False)
    return tb


def card(slide, x, y, w, h, fill=LIGHT):
    c = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    c.fill.solid(); c.fill.fore_color.rgb = fill; c.line.color.rgb = fill; c.shadow.inherit = False
    return c


def codebox(slide, x, y, w, h, code, size=11):
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = LIGHT; sp.line.color.rgb = LIGHT; sp.shadow.inherit = False
    tf = sp.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(12); tf.margin_right = Pt(6); tf.margin_top = Pt(8)
    for i, line in enumerate(code.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(1)
        r = p.add_run(); r.text = line if line else " "
        r.font.size = Pt(size); r.font.name = "Menlo"
        r.font.color.rgb = MUTED if line.strip().startswith("#") or line.strip().startswith("--") else INK
    return sp


def table(slide, rows, x, y, w, h, widths, size=12, head_size=13, first_col_accent=True, cell_fill=None):
    """rows[0] 是表頭。widths 為每欄 inch。cell_fill(r, c, text) 可覆寫底色。"""
    ncol = len(rows[0])
    tbl = slide.shapes.add_table(len(rows), ncol, x, y, w, h).table
    for i, cw in enumerate(widths):
        tbl.columns[i].width = Inches(cw)
    for r in range(len(rows)):
        for c in range(ncol):
            cell = tbl.cell(r, c)
            cell.margin_top = Pt(2); cell.margin_bottom = Pt(2); cell.margin_left = Pt(8); cell.margin_right = Pt(6)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            fill = None
            if cell_fill and r:
                fill = cell_fill(r, c, rows[r][c])
            cell.fill.fore_color.rgb = ACCENT if r == 0 else (fill or (WHITE if r % 2 else LIGHT))
            p = cell.text_frame.paragraphs[0]
            run = p.add_run(); run.text = rows[r][c]
            color = WHITE if r == 0 else (ACCENT if (c == 0 and first_col_accent) else INK)
            _font(run, size if r else head_size, color, r == 0 or c == 0)
    return tbl


def strip_rows(slide, items, y0=Inches(1.85), rh=Inches(1.0), gap=Inches(0.2), title_size=15, body_size=12.5):
    """左側色條 + 淺底卡片的列（面試框架 / 扣分點慣用）。items = [(title, body, strip_color)]"""
    cy = y0
    for title, body, strip in items:
        st = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), cy, Inches(0.12), rh)
        st.fill.solid(); st.fill.fore_color.rgb = strip; st.line.fill.background(); st.shadow.inherit = False
        bg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.72), cy, Inches(12.0), rh)
        bg.fill.solid(); bg.fill.fore_color.rgb = LIGHT; bg.line.fill.background(); bg.shadow.inherit = False
        textbox(slide, Inches(0.95), cy + Inches(0.1), Inches(11.6), rh - Inches(0.15),
                [(title, title_size, INK, True, 4), (body, body_size, MUTED, False)])
        cy += rh + gap




def flow(slide, steps, x0, y, bw, bh, gp, size=11):
    """水平步驟框 + 箭頭。steps = [(title, body, color)]"""
    for i, (t, d, color) in enumerate(steps):
        x = x0 + i * (bw + gp)
        box(slide, x, y, bw, bh, t + ("\n" + d if d else ""), color, WHITE, size)
        if i < len(steps) - 1:
            arrow(slide, x + bw, y + bh / 2, x + bw + gp, y + bh / 2)


def label(slide, x, y, w, text, size=10, color=MUTED, align=PP_ALIGN.LEFT, h=Inches(0.3)):
    return textbox(slide, x, y, w, h, [(text, size, color, False, 0)], align=align)


# ============ Slide 1: Title ============
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
bg.fill.solid(); bg.fill.fore_color.rgb = INK; bg.line.fill.background(); bg.shadow.inherit = False
box(s, Inches(0.9), Inches(2.0), Inches(1.5), Inches(1.5), "#n+1", ACCENT, WHITE, 32)
textbox(s, Inches(0.9), Inches(3.7), Inches(11.5), Inches(2.5), [
    ("Messenger", 44, WHITE, True, 6),
    ("System Design demo — 群聊即時收發、離線補收、同 chat 順序一致", 20, RGBColor(0xCB, 0xD5, 0xE1), False, 10),
    ("真多 process：2 gateway + 2 pub/sub shard owner + etcd + DynamoDB Local · kill 一台照樣連號無缺無重 · 附 AWS Terraform", 14, MUTED, False),
])

# ============ Slide 2: Requirements & scale ============
s = prs.slides.add_slide(BLANK)
header(s, "01 · REQUIREMENTS", "需求與規模")
cols = [
    ("功能性 (FR)", APP, [
        "建立群聊（≤ 100 人）",
        "收發訊息",
        "離線使用者回來後收到期間的訊息",
        "Out of scope：附件、視訊、線上狀態",
    ]),
    ("非功能性 (NFR)", DB, [
        "在線送達 < 500ms",
        "耐久：訊息不可遺失",
        "同一個 chat 所有人看到完全一致的順序",
        "1 億使用者、高吞吐",
    ]),
    ("容量數學", ANALYTICS, [
        "1 億條 WebSocket：Erlang 單機 100–200 萬 → 50–100 台；Python ~5 萬/process → 2,000 台",
        "每人每天 50 則 → 50 億則/天 ≈ 58K msg/s 平均、~175K 峰值",
        "fan-out 平均 10 人 → 175 萬投遞/s，但 owner 只推 M 個 gateway",
        "~200B × 50 億 ≈ 1TB/天",
        "S = 4096 shard；單 owner ~2 萬 msg/s → 峰值 ~9 台，留餘裕 ~20 台",
    ]),
]
cw = Inches(3.9); gap = Inches(0.25); x0 = Inches(0.6); y0 = Inches(1.8)
for i, (title, color, items) in enumerate(cols):
    x = x0 + i * (cw + gap)
    box(s, x, y0, cw, Inches(0.6), title, color, WHITE, 16)
    card(s, x, y0 + Inches(0.7), cw, Inches(4.4))
    bullets(s, x + Inches(0.2), y0 + Inches(0.9), cw - Inches(0.4), Inches(4.0), items, size=13, gap=10)
textbox(s, Inches(0.6), Inches(6.95), Inches(12.2), Inches(0.4), [
    ("設計成立、實作語言不成立：正式版 gateway 用 Go / Erlang / Rust（附錄 L）。DynamoDB 寫入峰值 ~70 萬 WCU，需事先申請 quota。", 12, MUTED, False)])

# ============ Slide 3: Protocols ============
s = prs.slides.add_slide(BLANK)
header(s, "02 · NETWORK PROTOCOLS", "協定比較：怎麼把訊息推到 client")
rows = [
    ("方式", "機制", "優點", "缺點", "適用"),
    ("Simple polling", "固定間隔 GET", "簡單", "大量空回應、高延遲", "偶爾檢查通知"),
    ("Long polling", "server hold 住直到有資料", "比 polling 即時", "每次回應後要重建請求；頻繁更新仍有延遲", "中低頻"),
    ("SSE", "server → client 單向串流", "簡單、自動重連", "單向；client 送訊仍要另一條 HTTP", "通知、價格推播"),
    ("WebSocket", "HTTP Upgrade 成持久雙向連線", "真雙向、低延遲", "連線管理與錯誤恢復較複雜", "聊天、遊戲、協作"),
]
table(s, rows, Inches(0.6), Inches(1.7), Inches(12.15), Inches(3.3), [1.9, 2.9, 1.9, 3.35, 2.1], size=12.5,
      cell_fill=lambda r, c, t: GREEN_BG if r == 4 else None)
box(s, Inches(0.6), Inches(5.35), Inches(12.15), Inches(1.3),
    "選 WebSocket + stateless gateway\n聊天是高頻雙向；gateway 只管連線、不擁有任何 chat → 使用者連任意一台、水平擴展容易。\n代價：連線是長期狀態 → 要處理斷線重連 + 補收（深入 3）。gRPC-Web 在瀏覽器要 proxy → 不選。", DB, WHITE, 13)

# ============ Slide 4: WebSocket API ============
s = prs.slides.add_slide(BLANK)
header(s, "03 · WEBSOCKET API", "WebSocket API（JSON frame，type 欄位）")
rows = [
    ("client → gateway", "欄位", "回應"),
    ("（連上 /ws?user_id=）", "正式版 ?token=<Cognito id token>", "welcome {gateway_id, chats[]} → 每個 chat 補收 newMessage{catchup} + catchupDone"),
    ("createChat（PDF）", "req_id, name, participants[]", "result {status SUCCESS|FAILURE, chat_id}"),
    ("sendMessage（PDF）", "req_id, chat_id, client_msg_id, text, client_ts", "result {status, message_id, dup}"),
    ("modifyChatParticipants（PDF）", "req_id, chat_id, user_id, operation ADD|REMOVE", "result {status, message_id(系統訊息)}"),
    ("fetchMessages", "req_id, chat_id, after_id, limit", "messages {items[], has_more}（跳號補拉 / 往回看）"),
    ("listChats", "req_id", "chats {chats[]}"),
    ("ack（PDF）", "chat_id, message_id（累積式）", "—"),
    ("ping", "t", "pong"),
]
table(s, rows, Inches(0.6), Inches(1.6), Inches(12.15), Inches(3.7), [3.0, 4.2, 4.95], size=11.5, head_size=12.5)
card(s, Inches(0.6), Inches(5.5), Inches(5.95), Inches(1.6))
textbox(s, Inches(0.8), Inches(5.6), Inches(5.6), Inches(1.5), [
    ("server 推送", 13, INK, True, 3),
    ("newMessage {chat_id, message_id, kind user|system, text, created_by, created_at, client_msg_id, client_ts, catchup} · chatAdded · chatRemoved · catchupDone · error", 11, MUTED, False)])
card(s, Inches(6.8), Inches(5.5), Inches(5.95), Inches(1.6))
textbox(s, Inches(7.0), Inches(5.6), Inches(5.6), Inches(1.5), [
    ("錯誤碼", 13, INK, True, 3),
    ("NOT_MEMBER · CHAT_NOT_FOUND · TOO_MANY_PARTICIPANTS · PAYLOAD_TOO_LARGE（> 4KB）· STORE_ERROR · UNAVAILABLE（重試 8 秒仍無 owner）· BAD_FRAME", 11, MUTED, False)])

# ============ Slide 5: Data model ============
s = prs.slides.add_slide(BLANK)
header(s, "04 · DATA MODEL", "資料模型（DynamoDB 5 張表）")
rows = [
    ("表", "鍵", "欄位 / 用途"),
    ("chat", "chat_id", "name, created_by, created_at；也是 createChat 的 commit marker"),
    ("membership", "chat_id PK + user_id SK", "查一個 chat 的所有成員（PDF）"),
    ("message", "chat_id PK + message_id SK", "text, created_by, kind, client_msg_id, client_ts, created_at"),
    ("chat_state", "user_id PK + chat_id SK", "last_read_message_id；PK=user_id 兼當「我有哪些 chat」索引"),
    ("client_msg", "chat_id PK + client_msg_id SK", "→ message_id；TTL 1 天；跨 owner 的持久去重"),
]
table(s, rows, Inches(0.6), Inches(1.6), Inches(12.15), Inches(2.4), [1.9, 3.4, 6.85], size=12.5)
notes = [
    ("PDF 的不一致", "ack 流程一處寫 Membership.last_read_message_id、一處寫 ChatState → 本專案統一放 chat_state；membership 只表達「誰在群裡」", AMBER),
    ("為什麼不用 GSI", "membership 上建 user_id GSI 也能查「我的 chat」，但 GSI 只有最終一致；chat_state 以 user_id 為 PK 可強一致讀。代價：createChat 多寫 N 筆", APP),
    ("交易上限 → commit-marker", "TransactWriteItems 上限 100 items，百人群需 201 筆 → 先 BatchWrite（冪等）、最後 conditional put chat 列 = commit 點", DB),
    ("新成員看不到加入前歷史", "ADD 時 chat_state 從目前 head 開始（PDF 未規定；多數聊天 app 的行為）", CDN),
]
strip_rows(s, notes, y0=Inches(4.2), rh=Inches(0.7), gap=Inches(0.05), title_size=12.5, body_size=10.5)

# ============ Slide 6: High-level architecture ============
s = prs.slides.add_slide(BLANK)
header(s, "05 · HIGH-LEVEL ARCHITECTURE", "整體架構")
y1 = Inches(2.0); y2 = Inches(3.4); bh = Inches(0.9)
box(s, Inches(0.5), y1, Inches(1.5), bh, "client\n(alice)", CLIENT, WHITE, 12)
box(s, Inches(0.5), y2, Inches(1.5), bh, "client\n(bob)", CLIENT, WHITE, 12)
box(s, Inches(2.5), y1, Inches(1.7), y2 + bh - y1, "nginx / ALB\n/ws least_conn\nUpgrade 轉發", GATEWAY, WHITE, 12)
arrow(s, Inches(2.0), y1 + bh / 2, Inches(2.5), y1 + bh / 2)
arrow(s, Inches(2.0), y2 + bh / 2, Inches(2.5), y2 + bh / 2)
label(s, Inches(1.95), y1 + bh / 2 - Inches(0.32), Inches(0.6), "WS", 10)
gx = Inches(4.7); gw = Inches(1.8)
box(s, gx, y1, gw, bh, "gateway-1\nstateless · 連線", APP, WHITE, 12)
box(s, gx, y2, gw, bh, "gateway-2\nstateless · 連線", APP, WHITE, 12)
arrow(s, Inches(4.2), y1 + bh / 2, gx, y1 + bh / 2)
arrow(s, Inches(4.2), y2 + bh / 2, gx, y2 + bh / 2)
px = Inches(8.6); pw = Inches(2.0)
box(s, px, y1, pw, bh, "pubsub-1\nshard 0,2,5…", ANALYTICS, WHITE, 12)
box(s, px, y2, pw, bh, "pubsub-2\nshard 1,3,4…", ANALYTICS, WHITE, 12)
for gy in (y1, y2):
    for py in (y1, y2):
        arrow(s, gx + gw, gy + bh / 2, px, py + bh / 2, color=MUTED, width=1.2)
textbox(s, Inches(6.5), Inches(1.45), Inches(2.1), Inches(0.5), [
    ("內部 WS（subscribe / publish / deliver）", 10, MUTED, False, 0)], align=PP_ALIGN.CENTER)
dx = Inches(11.1); dw = Inches(1.8)
box(s, dx, y1, dw, y2 + bh - y1, "DynamoDB\nchat · membership\nmessage\nchat_state\nclient_msg", DB, WHITE, 12)
arrow(s, px + pw, y1 + bh / 2, dx, y1 + bh / 2, color=DB)
arrow(s, px + pw, y2 + bh / 2, dx, y2 + bh / 2, color=DB)
# etcd：在 gateway 與 pubsub 欄之間的下方，線從方框下緣繞進來
ey = Inches(4.85); eh = Inches(0.8); ex = Inches(6.8); ew = Inches(1.6)
box(s, ex, ey, ew, eh, "etcd\nnodes/ · shards/", CDN, WHITE, 12)
gcx = gx + gw / 2; pcx = px + pw / 2; emid = ey + eh / 2
line(s, ex, emid, gcx, emid, color=CDN)
arrow(s, gcx, emid, gcx, y2 + bh, color=CDN)
label(s, gcx - Inches(1.9), emid - Inches(0.15), Inches(1.85), "watch → ShardMap", 10, CDN, PP_ALIGN.RIGHT)
line(s, pcx, y2 + bh, pcx, emid, color=CDN)
arrow(s, pcx, emid, ex + ew, emid, color=CDN)
label(s, pcx + Inches(0.08), emid - Inches(0.15), Inches(1.9), "lease + CAS claim", 10, CDN)
bullets(s, Inches(0.5), Inches(5.9), Inches(12.3), Inches(1.4), [
    "訊息流：sender → gateway → owner（定序 + 持久化）→ 訂閱該 chat 的 gateway → receiver；shard = crc32(topic) % S",
    "gateway 對每個 chat 只向 owner 訂閱一次（引用計數）→ fan-out 對 M 個 gateway，不是 N 個使用者",
    "dev 模式 ROLE=all：gateway + pub/sub 同一個 process（LocalLink）、SQLite、記憶體 registry",
], size=12, gap=4)

# ============ Slide 7: Decision map ============
s = prs.slides.add_slide(BLANK)
header(s, "決策地圖", "三大關鍵架構決策（環環相扣）")
cards = [
    ("① 怎麼把訊息推到 client", APP, "WebSocket\n+ stateless gateway", [
        "polling 大量空回應；long polling 每次重建請求",
        "聊天是高頻雙向 → WebSocket",
        "gateway 不擁有 chat，連任意一台都行",
        "代價：長連線狀態 → 斷線重連 + 補收",
    ]),
    ("② 跨 host 怎麼找到對方", DB, "chat 分片 pub/sub\n+ etcd registry", [
        "淘汰 user 分片：N 次 RPC、搬 user、無法定序",
        "shard = hash(chat_id) % S，唯一 owner",
        "gateway 向 owner 訂閱一次 → fan-out 對 M 台",
        "連鎖：需要 etcd lease + CAS 決定 owner",
    ]),
    ("③ 怎麼不亂、不丟", CACHE, "per-chat 連續序號\n+ 先持久化再 ack", [
        "per-chat lock 內 message_id = head + 1",
        "寫進 DB 成功才回 SUCCESS、才推送",
        "連續 → client 看到跳號就知道漏了、補拉",
        "連鎖：owner 會掛 → fencing + client_msg_id 去重",
    ]),
]
cw = Inches(3.9); gp = Inches(0.3); x0 = Inches(0.5); y = Inches(1.75)
for i, (title, color, choice, items) in enumerate(cards):
    x = x0 + i * (cw + gp)
    box(s, x, y, cw, Inches(0.6), title, color, WHITE, 15)
    box(s, x, y + Inches(0.68), cw, Inches(0.75), choice, GREEN, WHITE, 12)
    card(s, x, y + Inches(1.52), cw, Inches(3.1))
    bullets(s, x + Inches(0.2), y + Inches(1.68), cw - Inches(0.4), Inches(2.8), items, size=12, gap=8)
    if i < 2:
        arrow(s, x + cw, y + Inches(0.3), x + cw + gp, y + Inches(0.3))
textbox(s, Inches(0.5), Inches(6.5), Inches(12.3), Inches(0.6), [
    ("①決定「怎麼連」、②決定「怎麼找到對的人」、③決定「怎麼不亂、不丟」—— 下面逐項展開。", 13, MUTED, False)])

# ============ Slide 8: Deep dive 1 — user sharding vs chat sharding ============
s = prs.slides.add_slide(BLANK)
header(s, "深入 1 · ROUTING", "user 分片 vs chat 分片 pub/sub")
box(s, Inches(0.6), Inches(1.75), Inches(5.9), Inches(0.6), "（淘汰）user ownership consistent hashing", RED, WHITE, 15)
card(s, Inches(0.6), Inches(2.45), Inches(5.9), Inches(3.3))
bullets(s, Inches(0.8), Inches(2.6), Inches(5.5), Inches(3.1), [
    "registry 記 user_id → node；發訊時找每個收件人的 node 發 RPC",
    "問題 1：每則訊息 N 次 RPC（百人群聊 = 百次）",
    "問題 2：scale 時要搬 user ownership",
    "問題 3（致命）：同一個 chat 的訊息由不同 owner 處理 → 無法強制單一順序",
], size=13, gap=10)
box(s, Inches(6.85), Inches(1.75), Inches(5.9), Inches(0.6), "（本專案）以 chat 為核心的 pub/sub", GREEN, WHITE, 15)
card(s, Inches(6.85), Inches(2.45), Inches(5.9), Inches(3.3))
bullets(s, Inches(7.05), Inches(2.6), Inches(5.5), Inches(3.1), [
    "shard = crc32(topic) % S；每個 shard 同一時間只有一個 owner",
    "gateway 對每個 chat 向 owner subscribe(c:{chat}) 一次；本台第一個在看的人才訂、最後一個離開才退",
    "fan-out 對 M ≤ min(N, gateway 數) 個 gateway",
    "每個 chat 唯一 owner → 定序可行（深入 2）",
], size=13, gap=10)
arrow(s, Inches(6.5), Inches(2.05), Inches(6.85), Inches(2.05))
box(s, Inches(0.6), Inches(6.0), Inches(12.15), Inches(0.95),
    "不能用 Python hash()：字串 hash 每個 process 隨機化，兩台 gateway 會算出不同 shard → 用 crc32。\n固定 S 個 shard（同 Kafka partition / Redis Cluster slots）+ etcd 明確記 owner，比 hash ring 好推理（附錄 A）。", DB, WHITE, 13)

# ============ Slide 9: Deep dive 1 — etcd shard ownership ============
s = prs.slides.add_slide(BLANK)
header(s, "深入 1 · ETCD", "etcd key 佈局與 shard 擁有權")
codebox(s, Inches(0.6), Inches(1.65), Inches(5.6), Inches(2.6), """# pub/sub 節點（附 lease，TTL 5s）
nodes/{node_id}  → {ws}
# shard owner（附 owner 的 lease）
shards/{i}       → {node_id, ws}

# CAS claim：只有一人成功
txn compare create_revision(shards/i) == 0
    success put shards/i (lease=mine)""", size=12)
card(s, Inches(6.5), Inches(1.65), Inches(6.25), Inches(2.6))
bullets(s, Inches(6.7), Inches(1.75), Inches(5.9), Inches(2.5), [
    "lease grant(TTL 5s) → put nodes/{id} → 每 TTL/3 keepalive",
    "公平份額 ceil(S / 活著的節點數)；不足就搶沒人持有的",
    "新節點加入：持有過多者 drain → release（2 秒遲滯避免來回搶）",
    "節點死掉 → lease 過期 → key 全刪 → 他人 watch 到、jitter 後搶",
    "graceful（SIGTERM）主動 revoke lease → 立即轉移",
    "gateway：range 載入 + watch → 本地 ShardMap",
], size=12, gap=5)
textbox(s, Inches(0.6), Inches(4.45), Inches(12.2), Inches(0.4), [("實測（compose，S = 16）", 15, INK, True)])
states = [
    ("開機", "自動分成 8 / 8", [0, 1] * 8),
    ("hard kill pubsub-1", "~5.3 秒後 pubsub-2 接手全部 16", [1] * 16),
    ("pubsub-1 重新加入", "自動 rebalance 回 8 / 8", [1, 0] * 8),
]
cell = Inches(0.36); cg = Inches(0.06); gx0 = Inches(0.6); gy = Inches(4.95)
colw = Inches(4.1)
for k, (t, d, owners) in enumerate(states):
    bx = gx0 + k * colw
    for j, o in enumerate(owners):
        r, c = divmod(j, 8)
        box(s, bx + c * (cell + cg), gy + r * (cell + cg), cell, cell, str(j), ANALYTICS if o == 0 else APP, WHITE, 9,
            shape=MSO_SHAPE.RECTANGLE)
    textbox(s, bx, gy + 2 * (cell + cg) + Inches(0.05), Inches(3.4), Inches(0.7),
            [(t, 13, INK, True, 2), (d, 11.5, MUTED, False)])
    if k < 2:
        arrow(s, bx + 8 * (cell + cg), gy + cell, bx + colw - Inches(0.15), gy + cell)
label(s, Inches(0.6), Inches(7.0), Inches(12), "色塊 = owner：橘 pubsub-1、藍 pubsub-2（demo 頁 shard grid 同款）", 11)

# ============ Slide 10: Message flow sequence ============
s = prs.slides.add_slide(BLANK)
header(s, "訊息流 · SEQUENCE", "一則訊息的完整路徑")
actors = [("sender", CLIENT), ("gateway-1", APP), ("owner (pubsub)", ANALYTICS),
          ("DynamoDB", DB), ("gateway-2", APP), ("receiver", CLIENT)]
xs = [Inches(1.2), Inches(3.3), Inches(5.6), Inches(8.0), Inches(10.1), Inches(12.0)]
ty = Inches(1.55); aw = Inches(1.7); ah = Inches(0.5); bottom = Inches(6.35)
for (name, color), cx in zip(actors, xs):
    box(s, cx - aw / 2, ty, aw, ah, name, color, WHITE, 12)
    line(s, cx, ty + ah, cx, bottom, color=RGBColor(0xCB, 0xD5, 0xE1), width=1.0)


def msg(i, j, y, text, color=MUTED):
    x1, x2 = xs[i], xs[j]
    arrow(s, x1, y, x2, y, color=color, width=1.5)
    lw = max(abs(x2 - x1), Inches(2.7))
    label(s, (x1 + x2) / 2 - lw / 2, y - Inches(0.3), lw, text, 10.5, color, PP_ALIGN.CENTER)


msg(0, 1, Inches(2.5), "① sendMessage {client_msg_id}")
msg(1, 2, Inches(2.95), "② publish")
box(s, xs[2] - Inches(0.95), Inches(3.15), Inches(1.9), Inches(0.62), "③ async with chat.lock\nmid = head + 1", ANALYTICS, WHITE, 10.5)
msg(2, 3, Inches(4.1), "④ TransactWrite（2 items）", DB)
msg(3, 2, Inches(4.5), "OK：attribute_not_exists 通過", DB)
msg(2, 4, Inches(4.95), "⑤ deliver（仍在鎖內）", ANALYTICS)
msg(2, 1, Inches(5.3), "⑥ published {mid}", ANALYTICS)
msg(4, 5, Inches(5.3), "newMessage #mid")
msg(1, 0, Inches(5.75), "⑦ result SUCCESS", GREEN)
msg(5, 4, Inches(5.75), "ack（累積式）")
box(s, Inches(0.5), Inches(6.5), Inches(12.3), Inches(0.75),
    "SUCCESS 的時點 = DB 寫入成功之後：sender 收到 SUCCESS 時訊息一定已持久化；DB 失敗 → 不 fan-out、回 STORE_ERROR。\nfan-out 在鎖內 → 推送順序 = message_id 順序；gateway 每 500ms 批次把 ack 寫進 chat_state。", GREEN, WHITE, 12)

# ============ Slide 11: Deep dive 2 — sequential ids ============
s = prs.slides.add_slide(BLANK)
header(s, "深入 2 · ORDERING", "為什麼用 per-chat 連續序號")
rows = [
    ("方案", "有序", "唯一", "連續（可偵測漏訊）", "問題"),
    ("client timestamp", "✗", "✗", "✗", "clock skew；同毫秒不唯一"),
    ("Snowflake", "✓", "✓", "✗", "缺號是常態 → 無法判斷漏了沒"),
    ("UUIDv1", "✓", "✓", "✗", "同上"),
    ("per-chat 連續序號（owner 指派）", "✓", "✓", "✓", "單點指派：每 chat ~100 msg/s（百人上限足夠）"),
]
table(s, rows, Inches(0.6), Inches(1.65), Inches(12.15), Inches(2.4), [3.4, 1.0, 1.0, 2.3, 4.45], size=12.5,
      cell_fill=lambda r, c, t: GREEN_BG if r == 4 else (RED_BG if t == "✗" else None))
textbox(s, Inches(0.6), Inches(4.3), Inches(12.2), Inches(0.4), [("client 端排序規則（index.html onNewMessage）", 15, INK, True)])
rules = [
    ("id == lastSeen + 1", "顯示\n並 drain 緩衝", GREEN),
    ("id > lastSeen + 1", "漏了 → 緩衝\nfetchMessages(after=lastSeen)", AMBER),
    ("id ≤ lastSeen", "重複 → 丟掉", RED),
]
bw = Inches(3.9); gp = Inches(0.225)
for i, (t, d, color) in enumerate(rules):
    box(s, Inches(0.6) + i * (bw + gp), Inches(4.8), bw, Inches(1.1), t + "\n" + d, color, WHITE, 13)
textbox(s, Inches(0.6), Inches(6.1), Inches(12.2), Inches(1.0), [
    ("例：看到 #7、收到 #9 → 一定漏了 #8 → 先緩衝 #9、補拉 → 依序顯示。", 12.5, MUTED, False, 3),
    ("實測：兩個分頁分連 gw-1 / gw-2、同時各連發 20 則到同一個 chat → 雙方 #1–#44 順序完全一致、零重排零重複；跨 gateway ack p50 ~5ms。", 12.5, MUTED, False)])

# ============ Slide 12: Deep dive 3 — offline & durability ============
s = prs.slides.add_slide(BLANK)
header(s, "深入 3 · DURABILITY", "離線補收與耐久性")
cards = [
    ("persist-before-ack", DB, [
        "SUCCESS 時訊息一定已在 DynamoDB",
        "DB 寫失敗 → 不 fan-out、回 FAILURE STORE_ERROR",
        "client 保留「送出失敗、可重送」狀態",
    ]),
    ("client_msg_id 去重", APP, [
        "逾時重送、gateway 重試、跨 failover 重試都用同一個 id",
        "owner LRU（快）或 client_msg 表（跨 owner）",
        "→ 回同一個 message_id、dup=true",
    ]),
    ("先訂閱、再補收", ANALYTICS, [
        "連上 → welcome → subscribe → messages_after(last_read)",
        "反過來：補收完到訂閱前的訊息會漏掉",
        "先訂閱：最壞收到重複 → client 依 id 去重",
        "一次補收上限 200 則；has_more 時自己 fetch",
    ]),
    ("ack 合併", CDN, [
        "client 累積式 ack（只報最大連號）",
        "gateway 每 500ms 批次寫 chat_state（只進不退）",
        "gateway 掛 → 最多重送 500ms 內訊息",
        "at-least-once + client 去重 = 不遺失",
    ]),
]
cw = Inches(2.92); gp = Inches(0.17); x0 = Inches(0.6); y0 = Inches(1.75)
for i, (title, color, items) in enumerate(cards):
    x = x0 + i * (cw + gp)
    box(s, x, y0, cw, Inches(0.6), title, color, WHITE, 14)
    card(s, x, y0 + Inches(0.7), cw, Inches(3.3))
    bullets(s, x + Inches(0.15), y0 + Inches(0.85), cw - Inches(0.3), Inches(3.1), items, size=12, gap=8)
box(s, Inches(0.6), Inches(6.0), Inches(12.15), Inches(0.95),
    "實測：bob 斷線期間 alice 送 5 則 → bob 重連收到 #21–#25（catchup=true）；\n重送同 client_msg_id → dup=true、同一個 message_id。", GREEN, WHITE, 13)

# ============ Slide 13: Failover timeline ============
s = prs.slides.add_slide(BLANK)
header(s, "FAILOVER", "Failover 時序：hard vs graceful")
rows = [
    ("時間", "hard kill（SIGKILL pubsub-1，持有 chat X 的 shard 5）"),
    ("t = 0", "pubsub-1 被 SIGKILL"),
    ("t = 0 ~ 5s", "gateway → pubsub-1 內部 WS 斷線 → 清掉這些 topic 的訂閱；alice 送訊：link down → 等 shardmap 變更（重試，同 client_msg_id）"),
    ("t ≈ 5s", "etcd：pubsub-1 lease 過期 → nodes/pubsub-1、shards/* 被刪 → watch 事件"),
    ("t ≈ 5.1s", "pubsub-2 CAS 搶下 shard；狀態從 DB 重建（head = max(message_id)，consistent read）"),
    ("", "gateway watch 到新 owner → 重訂閱 → subscribed 回 head → 與 last_forwarded 比對，缺的從 DB 補送"),
    ("", "alice 那則重試到 pubsub-2 → 寫入 → SUCCESS（實測最慢一則 4.3s）"),
]
table(s, rows, Inches(0.6), Inches(1.6), Inches(8.3), Inches(4.1), [1.4, 6.9], size=11.5, head_size=12.5)
box(s, Inches(9.2), Inches(1.6), Inches(3.55), Inches(0.6), "hard kill", RED, WHITE, 15)
box(s, Inches(9.2), Inches(2.25), Inches(3.55), Inches(0.75), "≈ 5 秒", RED_BG, RED, 26)
label(s, Inches(9.2), Inches(3.05), Inches(3.55), "= lease TTL；最慢一則 4.3s", 11.5, MUTED, PP_ALIGN.CENTER)
box(s, Inches(9.2), Inches(3.65), Inches(3.55), Inches(0.6), "graceful（SIGTERM）", GREEN, WHITE, 15)
box(s, Inches(9.2), Inches(4.3), Inches(3.55), Inches(0.75), "42 ms", GREEN_BG, GREEN, 26)
label(s, Inches(9.2), Inches(5.1), Inches(3.55), "fence → revoke lease → 立即轉移", 11.5, MUTED, PP_ALIGN.CENTER)
bullets(s, Inches(0.6), Inches(5.95), Inches(12.2), Inches(1.3), [
    "500ms SLA 在 hard failover 窗口內不成立；降 TTL 可縮短，但 GC 停頓被誤判踢掉的機率上升",
    "gateway 最多重試 8 秒（涵蓋一次 hard failover），之後 FAILURE UNAVAILABLE，client 可用同 client_msg_id 重送",
], size=12.5, gap=5)

# ============ Slide 14: Fencing ============
s = prs.slides.add_slide(BLANK)
header(s, "FENCING", "Fencing 三層 + 殭屍舊 owner")
layers = [
    ("① etcd lease + CAS", CDN, "同一時間只有一個 owner\ncreate_revision == 0 才能 claim"),
    ("② 自我 fencing", AMBER, "keepalive 超過 TTL−1s 沒成功\n→ 拒絕 publish；SIGTERM 也先 fence"),
    ("③ DB conditional put", DB, "attribute_not_exists(message_id)\n撞到 → IdConflict"),
]
cw = Inches(3.9); gp = Inches(0.225)
for i, (t, color, d) in enumerate(layers):
    x = Inches(0.6) + i * (cw + gp)
    box(s, x, Inches(1.7), cw, Inches(0.55), t, color, WHITE, 15)
    box(s, x, Inches(2.3), cw, Inches(0.9), d, LIGHT, INK, 12, bold=False)
textbox(s, Inches(0.6), Inches(3.45), Inches(12.2), Inches(0.4), [("殭屍舊 owner（GC 停頓、網路分割）在 fencing 前寫了 N+1", 15, INK, True)])
steps = [
    ("新 owner 接手", "讀到 head = N", APP),
    ("殭屍寫入", "N+1 寫進 DB", RED),
    ("新 owner 寫 N+1", "attribute_not_exists\n→ IdConflict", AMBER),
    ("重讀 head", "向 etcd 確認\n仍是 owner", CDN),
    ("寫 N+2", "SUCCESS", DB),
    ("接收端", "看到跳號\n補拉 N+1", GREEN),
]
flow(s, steps, Inches(0.6), Inches(3.95), Inches(1.88), Inches(1.1), Inches(0.175), size=11.5)
box(s, Inches(0.6), Inches(5.3), Inches(12.15), Inches(0.6), "結果：無遺失、無重號 —— DB 順序就是唯一真相", GREEN, WHITE, 15)
bullets(s, Inches(0.6), Inches(6.05), Inches(12.2), Inches(1.2), [
    "逾時但其實寫成功：重試命中 client_msg 條件 → 同 id、dup=true",
    "etcd 整個掛掉：既有 owner 繼續服務（沒人能搶），但無法 failover → 正式版 etcd 3 節點跨 3 AZ",
], size=12.5, gap=5)

# ============ Slide 15: Verified results ============
s = prs.slides.add_slide(BLANK)
header(s, "實測 · VERIFIED", "實測結果（compose，本機）")
rows = [
    ("情境", "結果"),
    ("跨 gateway 連發 50 則", "全部 SUCCESS、1..50 連號、ack p50 5ms"),
    ("兩個分頁同時各連發 20 則到同一個 chat", "雙方 #1–#44 順序完全一致、零重排零重複"),
    ("hard kill chat owner（連發中）", "~5.3s 後存活節點接手 16 shard；60 則全 SUCCESS、連號無缺無重；最慢一則 4.3s"),
    ("graceful 下線 chat owner（連發中）", "60 則全 SUCCESS、最慢一則 42ms"),
    ("owner 重新加入", "自動 rebalance 回 8/8"),
    ("離線 5 則後重連", "收到 #21–#25（catchup）"),
]
table(s, rows, Inches(0.6), Inches(1.6), Inches(12.15), Inches(3.6), [4.4, 7.75], size=12.5, head_size=13,
      cell_fill=lambda r, c, t: GREEN_BG if c == 1 else None)
textbox(s, Inches(0.6), Inches(5.45), Inches(12.2), Inches(0.4), [("scripts/smoke_ws.py（dev 與 compose 共用）", 14, INK, True)])
codebox(s, Inches(0.6), Inches(5.85), Inches(12.15), Inches(1.35), """# 建群 → chatAdded → 連發 N 則（1..N 連號無重複）→ 離線補收 → 重送 dup → ADD 系統訊息 → 101 人被拒
python scripts/smoke_ws.py ws://localhost:8030/gw1/ws ws://localhost:8030/gw2/ws \\
       --burst 60 --interval 0.1 --kill-owner hard        # 送到 1/3 時殺掉這個 chat 的 owner""", size=11.5)

# ============ Slide 16: Local deployment ============
s = prs.slides.add_slide(BLANK)
header(s, "部署 · LOCAL", "本地部署（docker compose）")
rows = [
    ("服務", "port"),
    ("nginx（入口 + demo 頁；/ws least_conn、/gw1/ /gw2/ 指定）", "8030"),
    ("gateway-1 / gateway-2", "8031 / 8032"),
    ("pubsub-1 / pubsub-2（/debug/state）", "8041 / 8042"),
    ("etcd", "23790"),
    ("DynamoDB Local（init-tables 一次性建表）", "8101"),
]
table(s, rows, Inches(0.6), Inches(1.6), Inches(6.6), Inches(3.0), [5.0, 1.6], size=12)
codebox(s, Inches(7.5), Inches(1.6), Inches(5.25), Inches(1.4), """cd Messenger
docker compose up --build
open http://localhost:8030/
# 開兩個分頁：alice / bob""", size=12)
card(s, Inches(7.5), Inches(3.15), Inches(5.25), Inches(1.45))
textbox(s, Inches(7.7), Inches(3.22), Inches(4.9), Inches(1.35), [
    ("dev 模式（改 code 用）", 13, INK, True, 3),
    ("uvicorn app.main:app --port 8030 --timeout-graceful-shutdown 5；ROLE=all、SQLite、記憶體 registry，零外部依賴", 11, MUTED, False)])
textbox(s, Inches(0.6), Inches(4.85), Inches(12.2), Inches(0.4), [("推薦玩法", 15, INK, True)])
steps = [
    ("兩個分頁", "分連 gw-1 / gw-2", CLIENT),
    ("建群互傳", "看 #id 與 owner", APP),
    ("連發 20 則", "同時 hard kill\n白框 shard 的 owner", RED),
    ("~5 秒後", "shard grid 變色\n訊息照樣連號", ANALYTICS),
    ("compose start", "pubsub-1 重新加入\nrebalance 8/8", GREEN),
]
flow(s, steps, Inches(0.6), Inches(5.3), Inches(2.25), Inches(1.15), Inches(0.225), size=11.5)
label(s, Inches(0.6), Inches(6.7), Inches(12.2), "port 避開同 repo 其他 demo（8000/8010/8020、5432-5434、6379-6381）。", 11)

# ============ Slide 17: Prototype vs Production ============
s = prs.slides.add_slide(BLANK)
header(s, "PROTOTYPE → PRODUCTION", "原型 vs Production（env-gated）")
rows = [
    ("面向", "dev（單 process）", "本地部署 / 正式版", "切換 env"),
    ("程序角色", "ROLE=all：gateway + pub/sub 同一個 uvicorn", "ROLE=gateway / ROLE=pubsub 分開部署", "ROLE"),
    ("內部連線", "LocalLink（直接呼叫）", "WsLink（內部 WebSocket）", "依 ROLE"),
    ("Registry", "記憶體（本節點持有全部 shard）", "etcd", "ETCD_URL"),
    ("儲存", "SQLite", "DynamoDB Local（compose）/ DynamoDB（AWS）", "DYNAMO_TABLE_PREFIX · DYNAMO_ENDPOINT"),
    ("shard 數", "16", "16（compose）/ 64（AWS demo）/ 4096（正式）", "NUM_SHARDS"),
    ("入口", "uvicorn :8030", "nginx :8030 / CloudFront → ALB", "部署形態"),
    ("Auth", "?user_id=", "Cognito JWT ?token=", "AUTH_ENABLED"),
    ("debug kill", "開", "compose 開 / AWS 關", "ALLOW_DEBUG"),
]
table(s, rows, Inches(0.6), Inches(1.6), Inches(12.15), Inches(4.4), [1.8, 3.8, 3.85, 2.7], size=11.5, head_size=12.5)
textbox(s, Inches(0.6), Inches(6.35), Inches(12.2), Inches(0.8), [
    ("其他：LEASE_TTL_S（5）· MAX_PARTICIPANTS（100）· CATCHUP_LIMIT（200）· ACK_FLUSH_MS（500）· PUBLISH_DEADLINE_S（8）。本地部署 = 正式版 code path。", 12, MUTED, False)])

# ============ Slide 18: AWS deployment + CI/CD ============
s = prs.slides.add_slide(BLANK)
header(s, "部署 · AWS", "AWS 部署（Terraform）+ CI/CD")
yA = Inches(1.7); bh = Inches(0.85)
row = [
    (Inches(0.5), Inches(1.4), "Client\nwss://", CLIENT),
    (Inches(2.2), Inches(1.9), "CloudFront\n+ WAF", CDN),
    (Inches(4.4), Inches(2.1), "ALB\n只收 CloudFront\n+ X-Origin-Verify", GATEWAY),
    (Inches(6.8), Inches(1.9), "gateway ASG\nt4g.small ×2", APP),
    (Inches(9.0), Inches(1.9), "pubsub ASG\nt4g.small ×2", ANALYTICS),
]
for x, w, t, c in row:
    box(s, x, yA, w, bh, t, c, WHITE, 11)
for (x, w, *_), (nx, *_) in zip(row, row[1:]):
    arrow(s, x + w, yA + bh / 2, nx, yA + bh / 2)
label(s, Inches(8.35), yA - Inches(0.3), Inches(0.9), "內部 WS", 9.5, MUTED, PP_ALIGN.CENTER)
box(s, Inches(11.2), yA, Inches(1.6), bh, "DynamoDB\n5 張表", DB, WHITE, 11)
arrow(s, Inches(10.9), yA + bh / 2, Inches(11.2), yA + bh / 2, color=DB)
eY = Inches(2.95)
box(s, Inches(8.2), eY, Inches(1.6), Inches(0.6), "etcd t4g.micro", CDN, WHITE, 11)
line(s, Inches(8.2), eY + Inches(0.3), Inches(7.75), eY + Inches(0.3), color=CDN, width=1.2)   # etcd → gateway（watch）
arrow(s, Inches(7.75), eY + Inches(0.3), Inches(7.75), yA + bh, color=CDN, width=1.2)
line(s, Inches(10.4), yA + bh, Inches(10.4), eY + Inches(0.3), color=CDN, width=1.2)           # pubsub → etcd（lease + CAS）
arrow(s, Inches(10.4), eY + Inches(0.3), Inches(9.8), eY + Inches(0.3), color=CDN, width=1.2)
bullets(s, Inches(0.5), Inches(3.75), Inches(12.3), Inches(1.4), [
    "為何沒有 API Gateway：HTTP API 不能代理 WebSocket；WebSocket API 的 @connections fan-out = 每個連線一次 POST（PDF 淘汰的 N-RPC），且 2 小時連線上限、按訊息計費（附錄 D）",
    "idle timeout：ALB 300s、nginx 3600s；server 每 20s WS ping、client 每 25s app 層 ping。etcd demo 單台（SPOF），正式版 3 節點跨 3 AZ",
], size=11.5, gap=4)
textbox(s, Inches(0.5), Inches(5.0), Inches(12.2), Inches(0.4), [("CI/CD（deploy-messenger.yml，workflow_dispatch）", 14, INK, True)])
steps = [
    ("GitHub Actions", "OIDC assume role", GATEWAY),
    ("buildx", "linux/arm64\n→ ECR", APP),
    ("SSM 參數", "/messenger/\nIMAGE_URI", CDN),
    ("① pubsub 逐台", "max-concurrency 1\nSIGTERM → 轉移", ANALYTICS),
    ("② gateway", "一次 50%\nclient 重連另一半", APP),
    ("上線", "etcd 不受影響\n(tag 不同)", GREEN),
]
flow(s, steps, Inches(0.5), Inches(5.45), Inches(1.9), Inches(1.1), Inches(0.18), size=11)
label(s, Inches(0.5), Inches(6.75), Inches(12.2), "先 pubsub 逐台：同時間只有一台在換 → 不會所有 owner 同時下線；terraform apply 會計費，不用時 destroy。", 11)

# ============ Slide 19: Tech selection ============
s = prs.slides.add_slide(BLANK)
header(s, "技術選型", "被淘汰的選項與理由")
rows = [
    ("淘汰", "取代為", "理由"),
    ("Redis Pub/Sub", "自建 pub/sub（= shard owner）", "沒有持久化也不負責定序，仍要另做 sequencer；cluster 模式 7.0 前全叢集廣播"),
    ("Kafka（partition = shard）", "自建 pub/sub", "順序與耐久都好，但延遲與營運較重、gateway 要消費大量 partition；一 chat 一 topic 不可行"),
    ("Consul / ZooKeeper", "etcd", "Consul health check 模型偏服務、不是 lock；ZooKeeper JVM、ops 重"),
    ("Redis key + TTL registry", "etcd lease + txn", "不是共識系統，腦裂時會雙主"),
    ("Snowflake / UUIDv1", "per-chat 連續序號", "有序但不連續 → client 無法用跳號偵測漏訊息"),
    ("Cassandra", "DynamoDB", "LWT 慢、ops 重"),
    ("Postgres", "DynamoDB", "單 primary 寫入上限"),
    ("gRPC（內部）", "內部 WebSocket", "要 codegen；gRPC bidi 是正式版升級路徑"),
    ("API Gateway WebSocket API", "CloudFront → ALB", "@connections = N 次 RPC、2 小時連線上限、按訊息計費"),
]
table(s, rows, Inches(0.6), Inches(1.55), Inches(12.15), Inches(4.9), [2.9, 2.8, 6.45], size=11.5, head_size=12.5)
textbox(s, Inches(0.6), Inches(6.7), Inches(12.2), Inches(0.5), [
    ("etcd client：httpx 打 v3 JSON gateway，不裝 grpcio（aetcd 需 grpc；python-etcd3 已停更、與 protobuf 4 不相容）。", 12, MUTED, False)])

# ============ Slide 20: Interview framework + pitfalls ============
s = prs.slides.add_slide(BLANK)
header(s, "面試實戰 · FRAMEWORK", "45 分鐘怎麼講 + 常見扣分點")
iframe = [
    ("0–5 min：需求釐清", "群組上限？離線補收？順序要多嚴格？→ 算 1 億連線 / 58K–175K msg/s / 1TB/天", AMBER),
    ("5–15 min：協定 + API + 資料模型 + High-Level", "WebSocket vs polling/SSE；5 張表（chat_state 以 user_id 為 PK）；gateway / pub/sub / etcd / DynamoDB", ACCENT),
    ("15–35 min：三個深入（面試官挑 1–2 追問）", "① user 分片 vs chat 分片 + etcd｜② 連續序號 + client 跳號補拉｜③ persist-before-ack、先訂閱再補收", DB),
    ("35–45 min：failover 與收尾", "lease 過期 / graceful revoke、fencing 三層、殭屍 owner；Python 規模限制；反問面試官", GREEN),
]
strip_rows(s, iframe, y0=Inches(1.55), rh=Inches(0.72), gap=Inches(0.1), title_size=13, body_size=11)
textbox(s, Inches(0.6), Inches(4.85), Inches(12.2), Inches(0.4), [("常見扣分點", 15, RED, True)])
pitfalls = [
    ("用 client timestamp 定序", "clock skew、不唯一；要說 owner 指派 per-chat 序號"),
    ("以 user 分片", "N 次 RPC、搬 user、同 chat 多 owner → 無法定序"),
    ("ack 前沒持久化", "回了 SUCCESS 才寫 DB → owner 掛了訊息就丟"),
    ("離線補收空窗", "先補收再訂閱 → 中間訊息漏掉；要先訂閱、依 id 去重"),
    ("忽略 failover 雙主", "沒 lease / fencing → 殭屍 owner 重號；要說 CAS + conditional put"),
]
cw = Inches(2.35); gp = Inches(0.11); x0 = Inches(0.6); y = Inches(5.3)
for i, (t, d) in enumerate(pitfalls):
    x = x0 + i * (cw + gp)
    st = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(0.1), Inches(1.5))
    st.fill.solid(); st.fill.fore_color.rgb = RED; st.line.fill.background(); st.shadow.inherit = False
    card(s, x + Inches(0.1), y, cw - Inches(0.1), Inches(1.5))
    textbox(s, x + Inches(0.2), y + Inches(0.08), cw - Inches(0.3), Inches(1.4),
            [(t, 12.5, INK, True, 3), (d, 11, MUTED, False)])

# ============ Slide 21: Appendix overview ============
s = prs.slides.add_slide(BLANK)
header(s, "附錄總覽", "DESIGN.md 附錄 A–L 一頁列表")
apps = [
    ("A", "固定 S shard vs consistent hashing", "owner 狀態可從 DB 重建；要的是明確擁有權轉移；S 一開始給夠（4096）"),
    ("B", "Failover 與 fencing 時序", "lease 過期 ≈5s / graceful 42ms；逾時重試 dup、殭屍 owner、etcd 掛掉"),
    ("C", "DynamoDB 交易上限與 createChat", "201 items > 100 → BatchWrite + 最後 conditional put chat = commit 點"),
    ("D", "WebSocket 過 CloudFront / ALB / nginx", "淘汰 API GW；X-Origin-Verify；idle timeout 與 ping；部署斷線重連"),
    ("E", "Backpressure / 慢消費者", "每條連線送出佇列上限 1000，滿了就斷線（1013），重連走補收"),
    ("F", "100 人上限", "createChat 與 ADD 都檢查；ADD 在 owner 上與其他成員變更序列化"),
    ("G", "Rebalance", "公平份額下降 → drain（NOT_OWNER）→ release；gateway 刷新 shardmap 重試"),
    ("H", "WebSocket 上的 Auth", "瀏覽器 WS 不能自訂 header → ?token=；handshake 驗簽；log 遮 query"),
    ("I", "Monitoring", "CloudWatch + metric filter：shard_takeover / id_conflict / publish_failed"),
    ("J", "清理", "client_msg 用 DynamoDB TTL（1 天）；訊息永久保留；孤兒列可排程清"),
    ("K", "錯誤處理與降級", "owner 不在重試 8 秒 → UNAVAILABLE；DB 錯 STORE_ERROR；壞 frame 不斷線"),
    ("L", "Python 規模限制", "單 process ~5 萬 WS；架構可原封不動搬到 Go / Erlang"),
]
col_w = Inches(6.0); x0 = Inches(0.6); y0 = Inches(1.6); rh = Inches(0.88)
for i, (letter, title, desc) in enumerate(apps):
    col, row = (0, i) if i < 6 else (1, i - 6)
    x = x0 + col * (col_w + Inches(0.2)); y = y0 + row * rh
    box(s, x, y + Inches(0.05), Inches(0.5), Inches(0.6), letter, ACCENT, WHITE, 16)
    textbox(s, x + Inches(0.6), y, col_w - Inches(0.7), rh,
            [(title, 13, INK, True, 1), (desc, 11.5, MUTED, False)])

prs.save("Messenger.pptx")
print("saved Messenger.pptx ·", len(prs.slides._sldIdLst), "slides")
